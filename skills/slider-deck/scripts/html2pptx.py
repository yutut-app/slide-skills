#!/usr/bin/env python3
"""HTML を pptx に変換する。**HTML が編集対象、pptx はレビュー用の派生物。**

    python3 scripts/html2pptx.py deck.html -o deck/deck.pptx
    python3 scripts/html2pptx.py deck.html --template <テンプレ>.pptx -o out.pptx

テンプレートは HTML の <meta name="template" content="..."> から読む。
`--template` を渡せばそちらが優先。

**画像として貼らない。** 文字・表はネイティブの図形として作るので、
受け取った側が PowerPoint 上で数字を直せる。

**スライドマスターとレイアウトは書き換えない。** テンプレートを土台にして、
その上にプレースホルダ経由で中身を置く。スライドサイズもテンプレートに従う。

## 読み取る HTML の形

1スライド = <section class="slide" data-layout="...">。
data-layout は title / section / message-body / message-visual / table / blank。
省略すると中身から推測する。

    <section class="slide" data-layout="message-body">
      <h1>見出し（体言止め）</h1>
      <p class="message">メッセージライン（言い切り）</p>
      <ul><li>箇条書き<ul><li>入れ子は階層になる</li></ul></li></ul>
      <aside class="notes">スピーカーノート。スライドには出ない</aside>
    </section>

表は <table><thead><tr><th>…、図は <img src="...">、
出典は <p class="source">…</p>。

**class / style は見た目のためのもので、変換では見ない。**
位置は data-layout とテンプレートのレイアウトが決める。
HTML 側で細かく位置を指定しても pptx には反映されない。
"""

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from lxml import html as LH
from pptx import Presentation
from pptx.util import Emu, Inches, Pt

EMU_PER_INCH = 914400
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# ---------------------------------------------------------------------------
# テンプレートとレイアウト
# ---------------------------------------------------------------------------

def classify_layouts(prs):
    """レイアウトをプレースホルダの構成で分類する。名前では判定しない。

    レイアウト名はテンプレートごとに違い、日本語のこともある。
    """
    kinds = {k: [] for k in
             ("title", "section", "title_content", "title_only", "blank", "two_content")}
    for layout in prs.slide_layouts:
        names = [str(ph.placeholder_format.type) for ph in layout.placeholders]
        has_center = any("CENTER_TITLE" in n for n in names)
        has_title = any(n.startswith("TITLE") for n in names)
        has_sub = any("SUBTITLE" in n for n in names)
        content = [n for n in names if n.split()[0] in
                   ("OBJECT", "BODY", "PICTURE", "TABLE", "CHART")]
        if has_center:
            kinds["title"].append(layout)
        elif has_title and len(content) >= 2:
            kinds["two_content"].append(layout)
        elif has_title and len(content) == 1:
            kinds["title_content"].append(layout)
            if has_sub or "section" in layout.name.lower() or "セクション" in layout.name:
                kinds["section"].append(layout)
        elif has_title:
            kinds["title_only"].append(layout)
        else:
            kinds["blank"].append(layout)
    return kinds


FALLBACK = {
    "title": ["title", "title_content", "title_only", "blank"],
    "section": ["section", "title_only", "title_content", "blank"],
    "message-body": ["title_content", "title_only", "blank"],
    "message-visual": ["title_only", "title_content", "blank"],
    "table": ["title_only", "title_content", "blank"],
    "blank": ["blank", "title_only", "title_content"],
}


def pick_layout(kinds, want):
    for k in FALLBACK.get(want, FALLBACK["message-body"]):
        if kinds.get(k):
            return kinds[k][0], k
    return None, None


def blank_deck(template: Path):
    """テンプレートを土台にした、スライドが無いデッキ。テンプレート自体は触らない。"""
    import shutil
    import tempfile
    tmp = Path(tempfile.mkdtemp()) / "base.pptx"
    shutil.copy(template, tmp)
    prs = Presentation(str(tmp))
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.get(f"{R}id"))
        lst.remove(sld)
    return prs, tmp


# ---------------------------------------------------------------------------
# HTML を読む
# ---------------------------------------------------------------------------

def text_of(el):
    return " ".join(el.text_content().split()) if el is not None else ""


def read_list(ul, level=0, out=None):
    """<ul>/<ol> を (階層, 文言) の並びにする。入れ子は階層になる。"""
    out = [] if out is None else out
    for li in ul.findall("./li"):
        own = "".join(li.xpath("./text()")).strip()
        if not own:
            own = " ".join(
                t.strip() for t in li.xpath("./*[not(self::ul or self::ol)]//text()")
            ).strip()
        if own:
            out.append((level, " ".join(own.split())))
        for sub in li.findall("./ul") + li.findall("./ol"):
            read_list(sub, level + 1, out)
    return out


def read_table(tbl):
    header, rows = [], []
    for tr in tbl.xpath(".//tr"):
        cells = [text_of(c) for c in tr.xpath("./th|./td")]
        if not cells:
            continue
        if not header and tr.xpath("./th"):
            header = cells
        else:
            rows.append(cells)
    if not header and rows:
        header, rows = rows[0], rows[1:]
    return header, rows


def read_slides(doc):
    out = []
    for sec in doc.xpath("//section[contains(@class,'slide')]"):
        h = sec.xpath(".//h1|.//h2")
        msg = sec.xpath(".//*[contains(@class,'message')]")
        sub = sec.xpath(".//*[contains(@class,'subtitle')]")
        src = sec.xpath(".//*[contains(@class,'source')]")
        notes = sec.xpath(".//aside[contains(@class,'notes')]")
        lists = sec.xpath("./ul|./ol|.//div/ul|.//div/ol")
        tables = sec.xpath(".//table")
        imgs = sec.xpath(".//img")

        bullets = []
        for ul in lists:
            bullets += read_list(ul)

        layout = (sec.get("data-layout") or "").strip()
        if not layout:
            if tables:
                layout = "table"
            elif imgs:
                layout = "message-visual"
            elif bullets:
                layout = "message-body"
            elif sub:
                layout = "title"
            else:
                layout = "section" if h else "blank"

        out.append({
            "layout": layout,
            "title": text_of(h[0]) if h else "",
            "message": text_of(msg[0]) if msg else "",
            "subtitle": text_of(sub[0]) if sub else "",
            "bullets": bullets,
            "table": read_table(tables[0]) if tables else None,
            "image": imgs[0].get("src") if imgs else None,
            "source": text_of(src[0]) if src else "",
            "notes": text_of(notes[0]) if notes else "",
        })
    return out


# ---------------------------------------------------------------------------
# pptx を組む
# ---------------------------------------------------------------------------

def theme_fonts(prs):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    except KeyError:
        return None, None
    root = etree.fromstring(part.blob)
    node = root.find(f".//{A}fontScheme/{A}minorFont")
    if node is None:
        return None, None
    ea = node.find(f"{A}ea")
    latin = node.find(f"{A}latin")
    return (ea.get("typeface") if ea is not None else None,
            latin.get("typeface") if latin is not None else None)


def set_font(run, ea, latin):
    """**latin だけでなく ea にも書く。** 書かないと日本語に効かない。"""
    if latin:
        run.font.name = latin
    rPr = run.font._rPr
    for tag, face in (("ea", ea), ("cs", latin)):
        if not face:
            continue
        el = rPr.find(f"{A}{tag}")
        if el is None:
            el = rPr.makeelement(f"{A}{tag}", {})
            rPr.append(el)
        el.set("typeface", face)


def fill_ph(slide, kinds_startswith, value, ea, latin):
    for ph in slide.placeholders:
        name = str(ph.placeholder_format.type)
        if not name.startswith(kinds_startswith):
            continue
        tf = ph.text_frame
        tf.clear()
        items = value if isinstance(value, list) else [(0, value)]
        for i, (lvl, txt) in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.level = min(lvl, 4)
            r = p.add_run()
            r.text = txt
            set_font(r, ea, latin)
        return ph
    return None


def take_content_area(slide, prs):
    """本文プレースホルダの矩形を返し、その枠は取り除く。

    枠を残したまま図形を重ねると、空の枠線と「クリックしてテキストを入力」が
    残り、資料に写り込む。
    """
    for ph in list(slide.placeholders):
        name = str(ph.placeholder_format.type)
        if name.split()[0] in ("OBJECT", "BODY", "PICTURE", "TABLE", "CHART"):
            rect = (ph.left, ph.top, ph.width, ph.height)
            ph._element.getparent().remove(ph._element)
            return rect
    m = Emu(int(prs.slide_width * 0.06))
    top = Emu(int(prs.slide_height * 0.30))
    return (m, top, prs.slide_width - 2 * m, prs.slide_height - top - m)


def add_table(slide, area, header, rows, ea, latin):
    from pptx.util import Pt as _Pt
    ncols = max(len(header), max((len(r) for r in rows), default=0))
    nrows = len(rows) + (1 if header else 0)
    if not ncols or not nrows:
        return
    x, y, w, h = area
    gf = slide.shapes.add_table(nrows, ncols, x, y, w,
                                min(h, Inches(0.36) * nrows))
    tbl = gf.table
    data = ([header] if header else []) + rows
    for ri, row in enumerate(data):
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            cell.text = row[ci] if ci < len(row) else ""
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_font(r, ea, latin)
                    r.font.size = _Pt(11)


def add_image(slide, area, path: Path, base: Path):
    p = path if path.is_absolute() else (base / path)
    if not p.exists():
        return f"画像が見つからない: {p}"
    x, y, w, h = area
    pic = slide.shapes.add_picture(str(p), x, y, height=h)
    if pic.width > w:      # 幅がはみ出すなら幅基準に入れ直す
        slide.shapes._spTree.remove(pic._element)
        pic = slide.shapes.add_picture(str(p), x, y, width=w)
    pic.left = Emu(int(x + (w - pic.width) / 2))
    pic.top = Emu(int(y + (h - pic.height) / 2))
    return None


def add_note_text(slide, area, text, ea, latin, size=9):
    x, y, w, h = area
    tb = slide.shapes.add_textbox(x, Emu(int(y + h)), w, Inches(0.3))
    tb.text_frame.word_wrap = True
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    set_font(r, ea, latin)


def build(html_path: Path, template: Path, out: Path):
    doc = LH.parse(str(html_path)).getroot()
    slides = read_slides(doc)
    if not slides:
        sys.exit("スライドが見つからない。"
                 "<section class=\"slide\"> で1枚ずつ囲む必要がある。")

    prs, tmp = blank_deck(template)
    kinds = classify_layouts(prs)
    ea, latin = theme_fonts(prs)
    base = html_path.resolve().parent
    report, warn = [], []

    for i, s in enumerate(slides, 1):
        layout, got = pick_layout(kinds, s["layout"])
        if layout is None:
            warn.append(f"{i}枚目: 使えるレイアウトが無い")
            continue
        slide = prs.slides.add_slide(layout)

        if s["title"]:
            (fill_ph(slide, "TITLE", s["title"], ea, latin)
             or fill_ph(slide, "CENTER_TITLE", s["title"], ea, latin))

        body = []
        if s["message"]:
            body.append((0, s["message"]))
        body += [(lvl + (1 if s["message"] else 0), t) for lvl, t in s["bullets"]]

        if s["table"] or s["image"]:
            area = take_content_area(slide, prs)
            if s["message"]:
                # メッセージラインはタイトル直下に置く。図の上に重ねない
                x, y, w, h = area
                tb = slide.shapes.add_textbox(x, y, w, Inches(0.34))
                tb.text_frame.word_wrap = True
                r = tb.text_frame.paragraphs[0].add_run()
                r.text = s["message"]
                r.font.size = Pt(13)
                r.font.bold = True
                set_font(r, ea, latin)
                area = (x, Emu(int(y + Inches(0.46))), w,
                        Emu(int(h - Inches(0.46))))
            if s["table"]:
                add_table(slide, area, s["table"][0], s["table"][1], ea, latin)
            if s["image"]:
                err = add_image(slide, area, Path(s["image"]), base)
                if err:
                    warn.append(f"{i}枚目: {err}")
            if s["source"]:
                add_note_text(slide, area, s["source"], ea, latin)
        elif got == "title" and s["subtitle"]:
            fill_ph(slide, "SUBTITLE", s["subtitle"], ea, latin)
        elif body:
            (fill_ph(slide, "BODY", body, ea, latin)
             or fill_ph(slide, "OBJECT", body, ea, latin))

        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"]

        report.append((i, s["layout"], layout.name, got))

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    import shutil
    shutil.rmtree(tmp.parent, ignore_errors=True)

    print(f"変換: {html_path.name} → {out}")
    print(f"テンプレート: {template.name}  "
          f"{prs.slide_width / EMU_PER_INCH:.2f}×{prs.slide_height / EMU_PER_INCH:.2f}in  "
          f"書体: {ea or '（未指定）'}")
    print()
    print("| # | HTML の指定 | 使うレイアウト | 種別 |")
    print("|---|---|---|---|")
    for no, want, name, got in report:
        mark = "" if want == got or (want, got) in (("message-body", "title_content"),) else " ←妥協"
        print(f"| {no} | {want} | {name} | {got}{mark} |")
    if warn:
        print("\n**注意**")
        for w in warn:
            print(f"- {w}")
    print("\n**このあと必ず検査する。**")
    print("  python3 scripts/fit_check.py check " + str(out))
    print("  python3 scripts/qa_render.py " + str(out) + " -o qa")
    print("**直すのは HTML。出力された pptx を手で直さない。**再変換で消える。")
    return out


PRESENTATION_CT = ("application/vnd.openxmlformats-officedocument"
                   ".presentationml.presentation.main+xml")
TEMPLATE_CT = ("application/vnd.openxmlformats-officedocument"
               ".presentationml.template.main+xml")


def to_potx(pptx_path: Path, potx_path: Path):
    """pptx を potx（PowerPoint テンプレート）にする。

    **中身は変えない。** 変えるのは [Content_Types].xml の
    presentation パートの種別だけ。これが pptx と potx の違いのすべてで、
    スライド・レイアウト・マスターはそのまま使える。

    potx を PowerPoint で開くと、それを元にした新しいプレゼンテーションが
    作られる。テンプレートの見た目を実物で確かめるのに使う。
    """
    import shutil
    import zipfile

    src = zipfile.ZipFile(pptx_path)
    potx_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = str(potx_path) + ".tmp"
    changed = False
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                x = data.decode("utf-8")
                if PRESENTATION_CT in x:
                    x = x.replace(PRESENTATION_CT, TEMPLATE_CT)
                    changed = True
                data = x.encode("utf-8")
            out.writestr(item, data)
    src.close()
    shutil.move(tmp, potx_path)
    if not changed:
        print("注意: presentation の種別を書き換えられなかった。"
              "PowerPoint がテンプレートとして扱わない可能性がある")
    return potx_path


def expand_inputs(paths):
    """print.html のようなビューアを渡されたら、中の iframe の並びを使う。"""
    out = []
    for p in paths:
        path = Path(p)
        if not path.exists():
            sys.exit(f"見つからない: {path}")
        doc = LH.parse(str(path)).getroot()
        frames = doc.xpath("//iframe/@src")
        # iframe があればビューアとみなす。class の部分一致で判定すると
        # "slide-frame" が "slide" に当たって誤認する
        if frames:
            out += [(path.parent / f) for f in frames]
        else:
            out.append(path)
    return out


def guard_overwrite(out: Path, force: bool):
    """既存ファイルを黙って上書きしない。

    出力先に利用者が手で直したファイルが置かれていることがある。
    **上書きすると元に戻せない。** 別名を促す。
    """
    if not out.exists() or force:
        return
    import datetime
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M")
    alt = out.with_name(f"{out.stem}_{stamp}{out.suffix}")
    sys.exit(
        f"既にある: {out}\n"
        "**上書きしない。** 利用者が手で直したものかもしれない。\n"
        f"  別名で出す: -o {alt}\n"
        "  意図して置き換えるなら --force"
    )


def build_absolute(files, out: Path, px_per_pt=None, relayout=False):
    """px 絶対座標の HTML 群を1つの pptx にする。テンプレートの pptx は要らない。"""
    import html_abs as HA
    from pptx import Presentation as _P

    scale, src = HA.apply_scale_from(files, px_per_pt)

    prs = _P()
    w, h = HA.slide_size_px(LH.parse(str(files[0])).getroot())
    prs.slide_width = Emu(int(w * HA.EMU_PER_PX))
    prs.slide_height = Emu(int(h * HA.EMU_PER_PX))

    # 1枚目の body の書体をテーマにも書く。
    # テーマを直さないと、PowerPoint で新しく打った文字が既定の書体になる
    doc0 = LH.parse(str(files[0])).getroot()
    rules0 = []
    for s in doc0.xpath("//style"):
        rules0 += HA.parse_css(s.text or "")
    body0 = doc0.find(".//body")
    latin0, ja0 = HA.font_faces(HA.computed(body0, rules0) if body0 is not None else {})
    theme_ok = HA.set_theme_fonts(prs, latin0, ja0)

    warn, rows = [], []
    for i, f in enumerate(files, 1):
        doc = LH.parse(str(f)).getroot()
        before, n_before = len(warn), len(prs.slides._sldIdLst)
        HA.build_slide(prs, doc, f.resolve().parent, warn, relayout)
        made = len(prs.slides._sldIdLst) - n_before
        title = doc.findtext(".//title") or f.stem
        # **1ファイル=1枚とは限らない。**枚数を出さないと、
        # 15枚入りの HTML を渡して「1枚」で終わっても気づけない
        rows.append((i, f.name + (f"（{made}枚）" if made != 1 else ""),
                     title, len(warn) - before))

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))

    print(f"変換（絶対座標）: HTML {len(files)} ファイル → "
          f"**{len(prs.slides._sldIdLst)} 枚** → {out}")
    print(f"スライドサイズ: {w:.0f}×{h:.0f}px "
          f"= {w * HA.EMU_PER_PX / EMU_PER_INCH:.3f}×{h * HA.EMU_PER_PX / EMU_PER_INCH:.3f}in "
          f"= {w / scale:.0f}×{h / scale:.0f}pt")
    print(f"px→pt 換算: 1pt = {scale}px（出どころ: {src}）")
    print("**HTML の座標をそのまま写している。**テンプレートの pptx は使っていない")
    print(f"書体: 英数字={latin0 or '（未指定）'} / 日本語={ja0 or '（未指定）'}"
          f"（テーマにも{'書いた' if theme_ok else '書けなかった'}）\n")
    print("| # | ファイル | 種別 | 注意 |")
    print("|---|---|---|---|")
    for no, name, title, n in rows:
        print(f"| {no} | {name} | {title} | {n or ''} |")
    if warn:
        print("\n**注意**")
        for w_ in dict.fromkeys(warn):
            print(f"- {w_}")
    print("\n**このあと必ず検査する。**")
    print(f"  python3 scripts/fit_check.py check {out}")
    print(f"  python3 scripts/qa_render.py {out} -o qa")
    print("**直すのは HTML。出力された pptx を手で直さない。**")
    return out


def main():
    ap = argparse.ArgumentParser(
        description="HTML を pptx に変換する（HTML が編集対象、pptx は派生物）")
    ap.add_argument("html", nargs="+",
                    help="スライドの HTML。複数可。print.html を渡すと中の iframe 順に並べる")
    ap.add_argument("--template", "-t", help="テンプレートの .pptx。"
                    "省くと HTML の <meta name=\"template\"> を読む")
    ap.add_argument("-o", "--out", help="出力先。既定は <html名>.pptx")
    ap.add_argument("--force", action="store_true",
                    help="出力先が既にあっても上書きする。"
                         "**利用者が手で直したものでないと確かめてから使う**")
    ap.add_argument("--px-per-pt", type=float,
                    help="px→pt の換算係数。既定はマニフェストの `1pt = <n>px`、"
                         "それも無ければ 96dpi（1.3333）")
    ap.add_argument("--relayout", action="store_true",
                    help="**HTML の座標を写さない。**枠を行に束ね、余白をそろえて"
                         "敷き直す。器（ロゴ・上下の帯）は動かさない。"
                         "**重なりが出たら、まずこれを付けて変換し直す**")
    ap.add_argument("--potx", action="store_true",
                    help="PowerPoint テンプレート（.potx）として出す。"
                         "**実物で見た目を確かめるときに使う**")
    args = ap.parse_args()

    files = expand_inputs(args.html)

    # px 絶対座標のテンプレート（マニフェスト方式）なら、そちらの経路に入る
    import html_abs as HA
    if HA.looks_absolute(LH.parse(str(files[0])).getroot()):
        default_ext = ".potx" if args.potx else ".pptx"
        out = Path(args.out) if args.out else files[0].with_suffix(default_ext)
        pptx_out = out.with_suffix(".pptx") if args.potx else out
        guard_overwrite(out, args.force)
        build_absolute(files, pptx_out, args.px_per_pt, args.relayout)
        if args.potx:
            to_potx(pptx_out, out)
            if pptx_out != out:
                pptx_out.unlink(missing_ok=True)
            print(f"\nテンプレートとして出力: {out}")
            print("PowerPoint で開くと、これを元にした新しい資料が作られる。"
                  "**元 PPTX と並べて見比べる。**")
        return

    if len(files) > 1:
        sys.exit("レイアウト方式では HTML を1つずつ渡す。"
                 "1ファイルに <section class=\"slide\"> を並べる")
    html_path = files[0]

    tpl = args.template
    if not tpl:
        doc = LH.parse(str(html_path)).getroot()
        meta = doc.xpath("//meta[@name='template']/@content")
        tpl = meta[0] if meta else None
    if not tpl:
        sys.exit("テンプレートが指定されていない。\n"
                 "  HTML に <meta name=\"template\" content=\"...\"> を書くか、\n"
                 "  --template でテンプレートの .pptx を渡す。\n"
                 "**元の pptx は要らない。要るのはテンプレートだけ。**")
    template = Path(tpl)
    if not template.is_absolute():
        template = (html_path.resolve().parent / template).resolve()
    if not template.exists():
        sys.exit(f"テンプレートが見つからない: {template}")

    out = Path(args.out) if args.out else html_path.with_suffix(".pptx")
    guard_overwrite(out, args.force)
    build(html_path, template, out)


if __name__ == "__main__":
    main()
