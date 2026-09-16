#!/usr/bin/env python3
"""資料テンプレート（.pptx / .potx）から、フォントと文字サイズを読み出す。

図のテンプレートを資料テンプレートに合わせて作り直すために使う。
**フォントは本スキルで固定しない。資料テンプレートが決める。**

    python3 scripts/read_template_style.py assets/templates/deck/xxx.pptx

出力は次の3つ。
  1. テーマのフォント（見出し用・本文用、日本語と欧文それぞれ）
  2. レイアウトごとのプレースホルダと文字サイズ
  3. 実際に使われている色（テーマの配色）

読んだ結果は assets/template-map.md に書き写す。
そのうえで各生成スクリプトに --font を渡して図テンプレートを作り直す。
"""

import sys
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def theme_root(prs):
    """テーマ XML を解析して返す。

    テーマは汎用パートで _element を持たない。getattr で済ませると None が
    返るだけで、読めていないことに気づけない。blob を自分で解析する。
    """
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    except KeyError:
        return None
    return etree.fromstring(part.blob)


def theme_fonts(prs):
    """テーマの majorFont / minorFont を、latin と ea（日本語）に分けて返す。

    日本語フォントは ea（East Asian）に入る。latin だけ見ると英字用の指定しか
    取れず、日本語が何で表示されるか分からない。
    """
    out = {}
    root = theme_root(prs)
    if root is not None:
        scheme = root.find(f".//{A}fontScheme")
        if scheme is None:
            return out
        for kind in ("majorFont", "minorFont"):
            node = scheme.find(f"{A}{kind}")
            if node is None:
                continue
            latin = node.find(f"{A}latin")
            ea = node.find(f"{A}ea")
            out[kind] = {
                "latin": latin.get("typeface") if latin is not None else None,
                "ea": ea.get("typeface") if ea is not None else None,
            }
    return out


def layout_styles(prs):
    rows = []
    for i, layout in enumerate(prs.slide_layouts):
        for ph in layout.placeholders:
            sizes, fonts = [], []
            for para in ph.text_frame.paragraphs:
                if para.font.size:
                    sizes.append(para.font.size.pt)
                if para.font.name:
                    fonts.append(para.font.name)
                for run in para.runs:
                    if run.font.size:
                        sizes.append(run.font.size.pt)
                    if run.font.name:
                        fonts.append(run.font.name)
            rows.append({
                "layout_no": i,
                "layout": layout.name,
                "placeholder": ph.placeholder_format.type,
                "idx": ph.placeholder_format.idx,
                "sizes": sorted(set(sizes)),
                "fonts": sorted(set(fonts)),
            })
    return rows


def used_fonts(prs):
    """実際のスライドで使われているフォントを数える。テーマと食い違うことがある。"""
    c = Counter()
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        c[run.font.name] += 1
    return c


def theme_colors(prs):
    out = {}
    root = theme_root(prs)
    if root is not None:
        scheme = root.find(f".//{A}clrScheme")
        if scheme is None:
            return out
        for child in scheme:
            name = child.tag.split("}")[-1]
            srgb = child.find(f"{A}srgbClr")
            sysc = child.find(f"{A}sysClr")
            if srgb is not None:
                out[name] = srgb.get("val")
            elif sysc is not None:
                out[name] = sysc.get("lastClr")
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    if not src.exists():
        sys.exit(f"見つからない: {src}")

    prs = Presentation(str(src))

    print(f"# {src.name} のスタイル\n")
    print(f"スライドサイズ: {prs.slide_width.inches:.3f} × {prs.slide_height.inches:.3f} インチ\n")

    print("## テーマのフォント\n")
    tf = theme_fonts(prs)
    if tf:
        print("| 用途 | 欧文 | 日本語(ea) |")
        print("|---|---|---|")
        for kind, label in (("majorFont", "見出し"), ("minorFont", "本文")):
            v = tf.get(kind, {})
            print(f"| {label} | {v.get('latin') or '（未指定）'} | {v.get('ea') or '（未指定）'} |")
        ea = (tf.get("minorFont") or {}).get("ea")
        latin = (tf.get("minorFont") or {}).get("latin")
        print(f"\n**図テンプレートの再生成に使うフォント: `{ea or latin or '（要確認）'}`**")
        print("日本語(ea) が未指定なら、実際に使われているフォント（下）から決める。")
    else:
        print("テーマのフォントを読み取れなかった。実際に使われているフォント（下）から決める。")

    print("\n## 実際に使われているフォント（出現回数）\n")
    uf = used_fonts(prs)
    if uf:
        for name, n in uf.most_common():
            print(f"- {name}: {n}")
    else:
        print("（スライドに直接指定されたフォントは無い。テーマに従っている）")

    print("\n## レイアウトと文字サイズ\n")
    print("| # | レイアウト名 | プレースホルダ | idx | サイズ(pt) | フォント |")
    print("|---|---|---|---|---|---|")
    for r in layout_styles(prs):
        sizes = ", ".join(str(s) for s in r["sizes"]) or "（継承）"
        fonts = ", ".join(r["fonts"]) or "（継承）"
        print(f"| {r['layout_no']} | {r['layout']} | {r['placeholder']} | {r['idx']} | {sizes} | {fonts} |")

    print("\n## テーマの配色\n")
    tc = theme_colors(prs)
    if tc:
        print("| 役割 | 色 |")
        print("|---|---|")
        for k, v in tc.items():
            print(f"| {k} | `{v}` |")
        print("\n**配色もテンプレートに合わせる。** 本スキルの既定の4色は、"
              "テンプレートが無いときだけ使う。")
    else:
        print("（読み取れなかった）")

    print("\n---\n")
    print("この出力を assets/template-map.md に書き写し、"
          "次で図テンプレートを作り直す。\n")
    print("```bash")
    print("python3 scripts/gen_chart_templates.py --font '<日本語フォント名>'")
    print("python3 scripts/gen_qc7_templates.py   --font '<日本語フォント名>'")
    print("python3 scripts/gen_pptx_templates.py  --font '<日本語フォント名>'")
    print("```")


if __name__ == "__main__":
    main()
