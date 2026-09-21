#!/usr/bin/env python3
"""レイアウトの検算。**書く前に予算を出し、作った後に溢れを検出する。**

書く前（予算）:
    python3 scripts/fit_check.py budget --template <テンプレ>.pptx --layout 1
    python3 scripts/fit_check.py budget --slide-size 10.833x7.5 --box 9.0x3.2 --pt 16

作った後（検出）:
    python3 scripts/fit_check.py check <deck>.pptx

**生成してから直すのではなく、書く前に入る量を決める。**
往復の大半は「入らない量を書いてしまった」ことで起きる。

日本語の文字幅の見積もり:
    全角1文字の幅 ≒ フォントpt ÷ 72 インチ（等幅でない書体でも、全角はほぼ正方形）
    半角1文字の幅 ≒ その半分
    1行の高さ ≒ フォントpt × 行間倍率 ÷ 72 インチ

この見積もりは**目安であって保証ではない。** 書体によって1割程度ずれる。
最後は必ず画像にして目で見る（`qa_render.py`）。
"""

import argparse
import sys
import unicodedata
from pathlib import Path

EMU_PER_INCH = 914400
LINE_SPACING = 1.25      # 日本語の既定の行間倍率
SAFE_GAP_INCH = 0.15     # 要素どうしの最小の隙間
TEXT_MARGIN_INCH = 0.10  # テキスト枠の左右余白（片側）


# ---------------------------------------------------------------------------
# 文字幅の見積もり
# ---------------------------------------------------------------------------

def char_width_units(text: str) -> float:
    """文字列の幅を「全角何文字ぶん」で返す。半角は 0.5 と数える。"""
    total = 0.0
    for ch in text:
        total += 1.0 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 0.5
    return total


def max_chars_per_line(box_w_inch: float, pt: float,
                       margin_inch: float = TEXT_MARGIN_INCH) -> int:
    """1行に入る全角文字数。"""
    usable = box_w_inch - margin_inch * 2
    if usable <= 0 or pt <= 0:
        return 0
    return int(usable / (pt / 72))


def max_lines(box_h_inch: float, pt: float, spacing: float = LINE_SPACING) -> int:
    """枠に入る行数。"""
    if pt <= 0:
        return 0
    return int(box_h_inch / (pt * spacing / 72))


def needed_lines(text: str, box_w_inch: float, pt: float) -> int:
    """その文言が何行になるか。改行も数える。"""
    cols = max_chars_per_line(box_w_inch, pt)
    if cols <= 0:
        # **数えられないものを「溢れ」と言わない。**
        # 枠が文字1つぶんより狭いと 999 行必要と出し、
        # 1桁の数字（手順番号など）を溢れと誤検出していた
        return 0
    lines = 0
    for para in text.split("\n"):
        w = char_width_units(para)
        lines += max(1, -(-int(w * 100) // (cols * 100)))  # 切り上げ
    return lines


def table_height_inch(rows: int, row_pt: float, header_pt: float = None,
                      padding_inch: float = 0.06) -> float:
    """表の高さ。ヘッダー高 ＋ 行高 × 行数。"""
    header_pt = header_pt or row_pt
    header_h = header_pt * LINE_SPACING / 72 + padding_inch * 2
    row_h = row_pt * LINE_SPACING / 72 + padding_inch * 2
    return header_h + row_h * max(0, rows - 1)


# ---------------------------------------------------------------------------
# 予算を出す
# ---------------------------------------------------------------------------

def vertical_grid(slide_h: float, has_lead: bool = True):
    """縦方向のグリッドを先に固定する。

    **本文の上端と下端を先に決めてから中身を書く。**
    中身を書いてから場所を探すと、必ずどこかが重なる。
    比率で持つので、スライドサイズが変わっても崩れない。
    """
    g = {
        "title_top": slide_h * 0.043,
        "title_bottom": slide_h * 0.147,
        "lead_top": slide_h * 0.155,
        "lead_bottom": slide_h * 0.213,
        "body_top": slide_h * 0.245,
        "body_bottom": slide_h * 0.885,
        "footer_top": slide_h * 0.912,
    }
    if not has_lead:
        g["body_top"] = g["lead_top"]
    return g


def cmd_budget(args):
    if args.template:
        from pptx import Presentation
        prs = Presentation(args.template)
        sw = prs.slide_width / EMU_PER_INCH
        sh = prs.slide_height / EMU_PER_INCH
        print(f"テンプレート: {Path(args.template).name}")
        print(f"スライドサイズ: {sw:.3f} × {sh:.3f} インチ"
              f"（{'16:9' if abs(sw / sh - 16 / 9) < 0.02 else '4:3' if abs(sw / sh - 4 / 3) < 0.02 else '既定外'}）")
        print("**このサイズに合わせる。16:9 に直さない。** ブランドが崩れる\n")
        if args.layout is not None:
            layout = prs.slide_layouts[args.layout]
            print(f"レイアウト {args.layout}: {layout.name}")
            print("| プレースホルダ | idx | 幅(in) | 高さ(in) | 16pt での最大文字数/行 | 最大行数 |")
            print("|---|---|---|---|---|---|")
            for ph in layout.placeholders:
                w = (ph.width or 0) / EMU_PER_INCH
                h = (ph.height or 0) / EMU_PER_INCH
                print(f"| {ph.placeholder_format.type} | {ph.placeholder_format.idx} "
                      f"| {w:.2f} | {h:.2f} | {max_chars_per_line(w, 16)} | {max_lines(h, 16)} |")
            print()
    else:
        sw, sh = (float(x) for x in args.slide_size.lower().split("x"))
        print(f"スライドサイズ: {sw:.3f} × {sh:.3f} インチ\n")

    g = vertical_grid(sh)
    print("縦グリッド（先に固定する。中身を書いてから場所を探さない）")
    print("| 位置 | インチ |")
    print("|---|---|")
    for k, v in g.items():
        print(f"| {k} | {v:.2f} |")
    print(f"\n本文に使える高さ: {g['body_bottom'] - g['body_top']:.2f} インチ")
    print(f"要素どうしの最小の隙間: {SAFE_GAP_INCH} インチ\n")

    if args.box:
        bw, bh = (float(x) for x in args.box.lower().split("x"))
        print(f"枠 {bw} × {bh} インチ での文字数の上限")
        print("| フォント(pt) | 1行の最大文字数(全角) | 最大行数 | 合計(全角) |")
        print("|---|---|---|---|")
        for pt in ([args.pt] if args.pt else [11, 12, 14, 16, 18, 20]):
            c, l = max_chars_per_line(bw, pt), max_lines(bh, pt)
            print(f"| {pt} | {c} | {l} | {c * l} |")
        print("\n**上限を超える文言は書かない。** フォントを縮めて収めない。"
              "削るか、図にするか、スライドを分ける")


# ---------------------------------------------------------------------------
# 作った後に検出する
# ---------------------------------------------------------------------------

def rect(sh):
    try:
        if None in (sh.left, sh.top, sh.width, sh.height):
            return None
        return (sh.left / EMU_PER_INCH, sh.top / EMU_PER_INCH,
                sh.width / EMU_PER_INCH, sh.height / EMU_PER_INCH)
    except (AttributeError, ValueError):
        return None


def line_spacing_of(shp, pt, default=None):
    """**その図形に実際に書かれている行間を使う。** 既定値で決め打ちしない。

    行間は段落に書いてある。読まずに既定（1.35 など）で数えると、
    行間を詰めた枠を「溢れ」と誤判定する（実際に起きた）。
    倍率指定はそのまま、pt 固定指定は文字サイズで割って倍率に直す。
    """
    if default is None:
        default = LINE_SPACING
    if not shp.has_text_frame or pt <= 0:
        return default
    for para in shp.text_frame.paragraphs:
        v = para.line_spacing
        if v is None:
            continue
        return (v.pt / pt) if hasattr(v, "pt") else float(v)
    return default


def text_rect(shp, r, text, pt):
    """**枠ではなく、文字が実際に占める矩形を返す。**

    枠で重なりを判定すると誤検出が出る。全幅のタイトル枠と、右上の小さな
    ボックスは枠としては重なるが、**文字は左と右に分かれていて衝突しない。**
    縦中央揃えの隣り合う枠も、枠は重なるが文字は離れている。

    推定なので**枠より小さくは見積もるが、大きくはしない。**
    書体が分からない以上、狭く見積もりすぎると本物の重なりを見逃す。
    """
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    x, y, w, h = r
    lines = [ln for ln in text.split("\n") if ln.strip()] or [text]

    # 横: 最長行の推定幅。枠より広ければ枠のまま（折り返している）
    need_w = max(char_width_units(ln) for ln in lines) * pt / 72.0
    tf = shp.text_frame if shp.has_text_frame else None
    ml = (tf.margin_left or 0) / EMU_PER_INCH if tf else 0
    mr = (tf.margin_right or 0) / EMU_PER_INCH if tf else 0
    avail = max(w - ml - mr, 0.01)
    if need_w < avail:
        algn = None
        try:
            algn = tf.paragraphs[0].alignment if tf and tf.paragraphs else None
        except (AttributeError, IndexError):
            algn = None
        if algn == PP_ALIGN.CENTER:
            x = x + ml + (avail - need_w) / 2
        elif algn == PP_ALIGN.RIGHT:
            x = x + w - mr - need_w
        else:
            x = x + ml
        w = need_w

    # 縦: 行数ぶんの高さ。枠より高ければ枠のまま（溢れとして別に出る）
    need_h = len(lines) * pt * line_spacing_of(shp, pt) / 72.0
    if need_h < h:
        anchor = getattr(tf, "vertical_anchor", None) if tf else None
        if anchor == MSO_ANCHOR.MIDDLE:
            y = y + (h - need_h) / 2
        elif anchor == MSO_ANCHOR.BOTTOM:
            y = y + h - need_h
        h = need_h

    return (x, y, w, h)


def overlap_area(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    dx = min(ax + aw, bx + bw) - max(ax, bx)
    dy = min(ay + ah, by + bh) - max(ay, by)
    return dx * dy if dx > 0 and dy > 0 else 0.0


def body_pt(sh, default=18.0):
    """その図形の代表的な文字サイズ。指定が無ければ既定値。"""
    sizes = []
    if sh.has_text_frame:
        for para in sh.text_frame.paragraphs:
            if para.font.size:
                sizes.append(para.font.size.pt)
            for run in para.runs:
                if run.font.size:
                    sizes.append(run.font.size.pt)
    return max(sizes) if sizes else default


def segment(shp):
    """線（コネクタ）の端点を in で返す。取れなければ None。"""
    try:
        return (shp.begin_x / 914400, shp.begin_y / 914400,
                shp.end_x / 914400, shp.end_y / 914400)
    except (AttributeError, TypeError):
        return None


def crosses(r, seg) -> bool:
    """線分が長方形の内側を通るか（Liang–Barsky）。"""
    x, y, w, h = r
    x1, y1, x2, y2 = seg
    dx, dy = x2 - x1, y2 - y1
    t0, t1 = 0.0, 1.0
    for p, q in ((-dx, x1 - x), (dx, x + w - x1), (-dy, y1 - y), (dy, y + h - y1)):
        if p == 0:
            if q < 0:
                return False
            continue
        t = q / p
        if p < 0:
            if t > t1:
                return False
            t0 = max(t0, t)
        else:
            if t < t0:
                return False
            t1 = min(t1, t)
    return t0 <= t1


def filled(shp) -> bool:
    """塗りのある図形か。**帯やカードを見分けるため。**"""
    try:
        return shp.fill.type == 1        # MSO_FILL.SOLID
    except (AttributeError, TypeError, ValueError):
        return False


def centered(shp) -> bool:
    """段落が中央ぞろえか。1段落でも中央なら中央とみなす。"""
    from pptx.enum.text import PP_ALIGN
    if not shp.has_text_frame:
        return False
    return any(pa.alignment == PP_ALIGN.CENTER for pa in shp.text_frame.paragraphs)


def cmd_check(args):
    import checked
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(args.deck)
    sw = prs.slide_width / EMU_PER_INCH
    sh_ = prs.slide_height / EMU_PER_INCH
    findings = []
    gaps = {}   # 同じ位置の隙間をまとめる（枠の位置と大きさで引く）
    parts = {}  # 共通パーツ（画像）との重なり。同じ位置のものをまとめる
    n_slides, n_shapes = len(prs.slides._sldIdLst), 0

    for i, slide in enumerate(prs.slides, 1):
        items, lines = [], []
        for shp in slide.shapes:
            # コネクタと線は、図形に接するのが正しいので重なり判定から外す。
            # **ただし文字との重なりは別に見る**（下の「線と文字」）
            if shp.shape_type in (MSO_SHAPE_TYPE.LINE,):
                seg = segment(shp)
                if seg:
                    lines.append(seg)
                continue
            r = rect(shp)
            if not r:
                continue
            text = shp.text_frame.text.strip() if shp.has_text_frame else ""
            items.append((shp, r, text))
        n_shapes += len(items)

        # はみ出し
        for shp, r, _ in items:
            x, y, w, h = r
            if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + h > sh_ + 0.01:
                findings.append((i, "はみ出し",
                                 f"{shp.shape_type} が スライド外に出ている "
                                 f"({x:.2f},{y:.2f},{w:.2f}×{h:.2f} / 面 {sw:.2f}×{sh_:.2f})"))

        # 文字どうしの重なり
        for a in range(len(items)):
            for b in range(a + 1, len(items)):
                sa, ra, ta = items[a]
                sbb, rb, tb = items[b]
                if not (ta and tb):
                    continue  # 文字が無いものどうしの重なりは意匠のことがある
                # **枠ではなく文字の実占有域で判定する。**枠だと誤検出が出る
                ra_t = text_rect(sa, ra, ta, body_pt(sa))
                rb_t = text_rect(sbb, rb, tb, body_pt(sbb))
                ov = overlap_area(ra_t, rb_t)
                small = min(ra_t[2] * ra_t[3], rb_t[2] * rb_t[3])
                if small > 0 and ov / small > args.overlap_ratio:
                    findings.append((i, "重なり",
                                     f"「{ta[:14]}」と「{tb[:14]}」が "
                                     f"{ov / small * 100:.0f}% 重なっている"))

        # 溢れ（枠に入る行数を超えている）
        for shp, r, text in items:
            if not text:
                continue
            pt = body_pt(shp)
            # 折り返さない指定（white-space: nowrap 由来）の枠は、行数が増えない。
            # 折り返す前提で数えると、1行に収まる文言を溢れと誤判定する
            no_wrap = (shp.has_text_frame and shp.text_frame.word_wrap is False)
            if no_wrap:
                need = max(1, len([x for x in text.split("\n") if x.strip()]))
            else:
                need = needed_lines(text, r[2], pt)
            sp = line_spacing_of(shp, pt)
            cap = max_lines(r[3], pt, sp)
            if need == 0:
                # 幅から行数を数えられなかった。**溢れとも合格とも言わない**
                continue
            if cap and need > cap:
                findings.append((i, "溢れ",
                                 f"「{text[:14]}」が {need} 行必要だが枠は {cap} 行分"
                                 f"（{pt:.0f}pt・行間 {sp:.2f}・幅 {r[2]:.2f}in）"))

        # 文字が共通パーツ（ロゴなどの画像）に重なっている
        # **ロゴは画像なので、文字どうしの重なり判定には入らない。**タイトルがロゴに
        # かぶっても、はみ出し検査からも漏れていた。**同じ位置で全枚に出るものはまとめる**
        pics = [rect(sh) for sh in slide.shapes
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and rect(sh)]
        for sa, ra, ta in items:
            if not ta:
                continue
            tr = text_rect(sa, ra, ta, body_pt(sa))
            if not tr:
                continue
            for pr in pics:
                if overlap_area(tr, pr) > 0.02:
                    sig = (round(tr[0], 1), round(tr[1], 1), round(pr[0], 1), round(pr[1], 1))
                    parts.setdefault(sig, []).append((i, ta[:10]))
                    break

        # 線が文字を横切っている（軸ラベル・矢印・凡例と数値ラベル）
        # **意匠のこともあるので不合格にしない。**目で見る候補として出す。
        # 実績: 矢印のずれ・月の帯と文字・ラベルと線の重なりを、機械が1つも出さず見落とした
        for sa, ra, ta in items:
            if not ta:
                continue
            tr = text_rect(sa, ra, ta, body_pt(sa))
            if not tr:
                continue
            # 文字の実占有域を少し内側に取る。**端をかすめる線は拾わない**
            x, y, w, h = tr
            inner = (x + w * 0.08, y + h * 0.2, w * 0.84, h * 0.6)
            for (x1, y1, x2, y2) in lines:
                if crosses(inner, (x1, y1, x2, y2)):
                    findings.append((i, "線と文字",
                                     f"線が「{ta[:10]}」の文字に重なっている"
                                     f"（({x1:.2f},{y1:.2f})→({x2:.2f},{y2:.2f})）。**目で確かめる**"))
                    break

        # 帯の上の文字が、帯の中心とそろっていない
        # **意匠の話なので不合格にしない。配布元に報告するだけ。**
        # 帯（塗りがあって文字を持たない枠）に、中央ぞろえの文字が載っていて、
        # 文字の枠が帯の 6 割以上を占めるなら、**帯いっぱいに見せる意図**とみなす。
        # バッジのように意図して片側へ寄せたものを拾わないための線引き。
        for sa, ra, ta in items:
            if ta or not filled(sa):
                continue
            for sb, rb, tb in items:
                if sb is sa or not tb or not centered(sb):
                    continue
                if not (ra[0] <= rb[0] and rb[0] + rb[2] <= ra[0] + ra[2]
                        and ra[1] - 0.03 <= rb[1] <= ra[1] + ra[3] + 0.03):
                    continue
                if ra[2] <= 0 or rb[2] < ra[2] * 0.6:
                    continue
                dx = (rb[0] + rb[2] / 2) - (ra[0] + ra[2] / 2)
                if abs(dx) > 0.02:
                    findings.append((i, "中心ずれ",
                                     f"「{tb[:10]}」が、下の帯の中心から {dx:+.2f}in ずれている"
                                     f"（帯 {ra[2]:.2f}in / 文字の枠 {rb[2]:.2f}in）。"
                                     f"**帯と同じ left/width にする。配布元の担当**"))

        # 隙間が足りない（縦に並ぶ要素）。**まとめてから足す**（下の gaps）
        text_items = sorted([(r, t, body_pt(s)) for s, r, t in items if t],
                            key=lambda x: x[0][1])
        for a in range(len(text_items) - 1):
            (x1, y1, w1, h1), t1, pt1 = text_items[a]
            (x2, y2, w2, h2), t2, pt2 = text_items[a + 1]
            if overlap_area((x1, y1, w1, h1), (x2, y2, w2, h2)) > 0:
                continue  # 重なりとして別に報告済み
            if min(x1 + w1, x2 + w2) - max(x1, x2) <= 0:
                continue  # 横に並んでいるので縦の隙間は関係ない
            gap = y2 - (y1 + h1)
            if not (0 <= gap < SAFE_GAP_INCH):
                continue
            # 表の行や積み上げた枠は、接しているのが正しい。
            # 幅と左端がそろっていて隙間がほぼ0なら、1つの構造とみなす
            if gap < 0.02 and abs(x1 - x2) < 0.02 and abs(w1 - w2) < 0.02:
                continue
            # 見出しとその説明は1つのまとまり。近づけるのが正しい（近接の原則）。
            # 上が明らかに大きい字で、左端がそろっているなら同じブロックとみなす
            if pt1 >= pt2 * 1.3 and abs(x1 - x2) < 0.06:
                continue
            # **同じ場所の隙間は、枚ごとに数えない。**
            # テンプレートが決めている間隔は全枚で同じなので、
            # 1枚ずつ出すと本物の指摘が埋もれる（18枚で27件出た）。
            # 文字は枚ごとに違うので、**枠の位置と大きさでまとめる**
            sig = (round(x1, 2), round(w1, 2), round(h1, 2),
                   round(x2, 2), round(w2, 2), round(gap, 2))
            gaps.setdefault(sig, []).append(
                (i, f"「{t1[:10]}」と「{t2[:10]}」", gap))

    for sig, hits in parts.items():
        slides_ = sorted({h[0] for h in hits})
        if len(slides_) >= 3:
            findings.append((slides_[0], "共通パーツ",
                             f"「{hits[0][1]}」がロゴなどの画像に重なっている"
                             f"（**{len(slides_)}枚で同じ位置**: {slides_[0]}〜{slides_[-1]}枚目）。"
                             "**テンプレートの配置。配布元の担当**"))
        else:
            for i_, t in hits:
                findings.append((i_, "共通パーツ",
                                 f"「{t}」がロゴなどの画像に重なっている。**目で確かめる**"))

    # **同じ場所のものは1件にまとめる。**枚数を添えて、どこの話か分かるようにする
    for sig, hits in sorted(gaps.items(), key=lambda kv: -len(kv[1])):
        gap = sig[5]
        slides_ = sorted({h[0] for h in hits})
        if len(slides_) >= 3:
            findings.append((slides_[0], "隙間不足",
                             f"左 {sig[0]:.1f}in・上下の幅 {sig[1]:.1f}/{sig[4]:.1f}in "
                             f"の隙間が {gap:.2f}in "
                             f"（**{len(slides_)}枚で同じ**: {slides_[0]}〜{slides_[-1]}枚目）。"
                             f"**テンプレートが決めている間隔。配布元の担当**"
                             f"（{SAFE_GAP_INCH}in 以上あける）"))
        else:
            for i_, pair, g in hits:
                findings.append((i_, "隙間不足",
                                 f"{pair}の間が {g:.2f}in"
                                 f"（{SAFE_GAP_INCH}in 以上あける）"))

    print(f"検査: {Path(args.deck).name}  {len(prs.slides)} 枚  "
          f"{sw:.2f}×{sh_:.2f}in")
    if not findings:
        print("\n重なり・はみ出し・溢れ・隙間不足は見つからなかった。")
        print("**これで終わりにしない。** 数値の見積もりは目安であり、"
              "書体によってずれる。必ず qa_render.py で画像にして目で見る。")
        # **ここでもサマリを出す。** 指摘ゼロの経路が素通りすると、
        # 「0件」と「何も見ていない」が区別できない（実際にそうなっていた）
        return checked.summary("fit_check", n_shapes, "図形", 0, {
            "スライド": f"{n_slides}枚",
            "行間": "図形に書かれた値",
            "書体": checked.font_state("verdana", "meiryo"),
        })

    # 重大度を3つに分ける。
    #   不合格 … 座標だけで決まる客観的な欠陥。書体が無くても正しく出る
    #   要確認 … **書体が無いときの「溢れ」。**文字幅が見積もりなので当てにならない
    #   注意   … 設計の余地がある
    #
    # **書体が無いときの溢れを不合格に混ぜない。**
    # 実案件で「毎回これは書体起因だと説明する手間」が生じた。
    # 判定できないものを不合格と呼ぶと、本物の不合格が埋もれる。
    FATAL = ("重なり", "はみ出し", "溢れ")
    fonts_ok = checked.font_state("verdana", "meiryo") == "そろい"
    fatal = [f for f in findings
             if f[1] in FATAL and (fonts_ok or f[1] != "溢れ")]
    unsure = [] if fonts_ok else [f for f in findings if f[1] == "溢れ"]
    warn = [f for f in findings if f[1] not in FATAL]

    if fatal:
        print(f"\n## 不合格: {len(fatal)} 件（重なり・はみ出し・溢れ）\n")
        print("| スライド | 種別 | 内容 |")
        print("|---|---|---|")
        for no, kind, msg in fatal:
            print(f"| {no} | {kind} | {msg} |")
        print("\n**ゼロになるまで直す。回数で打ち切らない。**")
        print("直し方はフォントを縮めるのではなく、"
              "**内容を削るか、図にするか、スライドを分ける**。")

    if unsure:
        print(f"\n## 要確認: {len(unsure)} 件（溢れ／**書体がそろうまで判定できない**）\n")
        print("| スライド | 種別 | 内容 |")
        print("|---|---|---|")
        for no, kind, msg in unsure:
            print(f"| {no} | {kind} | {msg} |")
        print("\n**指定の書体が手元にそろっていない。** 文字幅を "
              "全角=pt/72in の見積もりで数えているので、**実際より多く出る。**")
        print("**不合格には数えない。**画像で見て収まっていれば、その枚は問題ない。")
        print("**画像を見ずに文言を削らない。**書体がそろった手元で数え直す。")
    if warn:
        print(f"\n## 注意: {len(warn)} 件（設計の余地あり。自動では不合格にしない）\n")
        print("| スライド | 種別 | 内容 |")
        print("|---|---|---|")
        for no, kind, msg in warn:
            print(f"| {no} | {kind} | {msg} |")
        print("\n見た目の好みの調整はここ。**この一覧を空にすることを目標にしない。**")

    if not findings:
        pass
    print("\n**画像にして目で見るまで終わりにしない**（qa_render.py）。"
          "文字幅の見積もりは目安で、書体によってずれる。")

    rc = checked.summary("fit_check", n_shapes, "図形", len(findings), {
        "スライド": f"{n_slides}枚",
        "行間": "図形に書かれた値",
        "文字幅": "全角=pt/72inの見積もり",
        "書体": checked.font_state("verdana", "meiryo"),
        "不合格": len(fatal),
        "要確認": len(unsure),
    })
    # **要確認では落とさない。**判定できないものを不合格と同じ扱いにしない
    return rc or (1 if fatal else 0)


def main():
    ap = argparse.ArgumentParser(description="レイアウトの検算。budget は書く前、check は変換の直後")
    sub = ap.add_subparsers(dest="cmd", required=True)

    b = sub.add_parser("budget", help="書く前に、入る文字数と縦グリッドを出す")
    b.add_argument("--template", help="テンプレートの .pptx。サイズをここから取る")
    b.add_argument("--layout", type=int, help="レイアウト番号。枠の寸法を出す")
    b.add_argument("--slide-size", default="13.333x7.5",
                   help="テンプレートが無いときのサイズ。例 10.833x7.5")
    b.add_argument("--box", help="文字数を計算したい枠。例 9.0x3.2")
    b.add_argument("--pt", type=float, help="フォントサイズ。省くと代表値を一覧で出す")
    b.set_defaults(func=cmd_budget)

    c = sub.add_parser("check", help="作った後に、重なり・はみ出し・溢れを検出する")
    c.add_argument("deck")
    c.add_argument("--overlap-ratio", type=float, default=0.15,
                   help="重なりとみなす面積比（既定 0.15）")
    c.set_defaults(func=cmd_check)

    args = ap.parse_args()
    sys.exit(args.func(args) or 0)


if __name__ == "__main__":
    main()
