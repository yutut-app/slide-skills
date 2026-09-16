#!/usr/bin/env python3
"""作図系テンプレート（PowerPoint）を生成する。

1テンプレ = 1ファイル = 1スライド（16:9）。
図形はすべて編集可能なオートシェイプで作る。画像にしない。

道具の名前と用途は mfg-improvement-frameworks スキルの
references/30_qc7.md・40_n7.md に揃える。勝手に言い換えない。

使い方と注意はスライド上ではなくスピーカーノートに入れる。
スライドに注意書きを載せると、そのまま資料に貼られてしまうため。

使い方:
    python3 scripts/gen_pptx_templates.py [qc7出力先] [n7出力先]
    既定は assets/templates/qc7 と assets/templates/n7
"""

import sys
import pathlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# 提案・報告資料向けの抑えた配色。公式 pptx スキルの派手なパレットは使わない
NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x44, 0x72, 0xC4)
LIGHT = RGBColor(0xD9, 0xE2, 0xF3)
PALE = RGBColor(0xF2, 0xF5, 0xFA)
GRAY = RGBColor(0x80, 0x80, 0x80)
LINE_GRAY = RGBColor(0xBF, 0xBF, 0xBF)
ACCENT = RGBColor(0xC0, 0x00, 0x00)
INK = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "游ゴシック"        # 日本語（ea）
LATIN = None               # 欧文（latin）。None なら日本語と同じものを使う
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def apply_font(run, size=None, bold=None, color=None):
    """書体を run に当てる。**latin だけでなく ea と cs にも書く。**

    python-pptx の font.name は <a:latin> しか書かない。テーマの ea が空だと、
    日本語は script 別フォント（既定はＭＳ Ｐゴシック）に落ち、指定した書体が
    効かない。**日本語を扱う以上、ea を自分で書く。**
    """
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color is not None:
        f.color.rgb = color
    f.name = LATIN or FONT                      # <a:latin>
    rPr = f._rPr
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for tag, face in (("ea", FONT), ("cs", LATIN or FONT)):
        el = rPr.find(f"{A}{tag}")
        if el is None:
            el = rPr.makeelement(f"{A}{tag}", {})
            rPr.append(el)
        el.set("typeface", face)
    return run


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def textbox(slide, x, y, w, h, text, size=11, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    apply_font(r, size=size, bold=bold, color=color)
    return tb


def _drop_theme_style(sp):
    """オートシェイプ既定の <p:style> を外す。

    これを残すとテーマの effectRef が効き、影が付く。影は資料を古く見せるので消す。
    """
    el = sp._element
    style = el.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        el.remove(style)


def _fill_text(tf, text, size, bold, color, align=PP_ALIGN.CENTER):
    """改行を含む文字列を段落に分けて流し込む。

    "a\\nb" を1つの run に入れると、2行目だけ配置がずれる。
    改行は必ず段落として扱う。
    """
    lines = str(text).split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        apply_font(r, size=size, bold=bold, color=color)


def shape(slide, kind, x, y, w, h, text="", fill=LIGHT, line=NAVY,
          size=11, bold=False, color=INK, line_w=1.0):
    sp = slide.shapes.add_shape(kind, x, y, w, h)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    _drop_theme_style(sp)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _fill_text(tf, text, size, bold, color)
    return sp


def edge_point(rect, toward):
    """矩形 rect の中心から toward の中心へ向かう線が、rect の辺と交わる点。

    矢印の始点・終点をこの点にしないと、線が箱の内側まで入り込んで
    文字に重なる。連関図のように放射状に線を引く図で必ず要る。
    rect / toward は (x, y, w, h) の EMU タプル。
    """
    x, y, w, h = rect
    cx, cy = x + w / 2, y + h / 2
    tx, ty, tw, th = toward
    dx, dy = (tx + tw / 2) - cx, (ty + th / 2) - cy
    if dx == 0 and dy == 0:
        return cx, cy
    # 縦横それぞれで辺に到達する倍率を求め、小さい方（先に当たる辺）を採る
    scales = []
    if dx != 0:
        scales.append((w / 2) / abs(dx))
    if dy != 0:
        scales.append((h / 2) / abs(dy))
    t = min(scales)
    return Emu(int(cx + dx * t)), Emu(int(cy + dy * t))


def rect_of(sp):
    return (sp.left, sp.top, sp.width, sp.height)


def box(slide, x, y, w, h, text="", **kw):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, text, **kw)


def rbox(slide, x, y, w, h, text="", **kw):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, text, **kw)


def line(slide, x1, y1, x2, y2, color=NAVY, width=1.25, dash=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    if dash:
        cn.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    return cn


def arrow(slide, x1, y1, x2, y2, color=NAVY, width=1.25):
    """矢印。python-pptx はコネクタに矢尻を直接設定できないので XML を触る。"""
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color
    cn.line.width = Pt(width)
    ln = cn.line._get_or_add_ln()
    tail = ln.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
        {"type": "triangle", "w": "med", "len": "med"},
    )
    ln.append(tail)
    return cn


def frame(slide, title, message):
    """スライド共通の枠。左上にタイトル、その下にメッセージライン。

    見出しは体言止め、メッセージラインは言い切り。句点は付けない。
    """
    textbox(slide, Inches(0.5), Inches(0.32), Inches(12.4), Inches(0.4),
            title, size=20, bold=True, color=NAVY)
    textbox(slide, Inches(0.5), Inches(0.82), Inches(12.4), Inches(0.34),
            message, size=13, bold=True, color=INK)
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.26), Inches(12.83), Inches(1.26))
    ln.line.color.rgb = LINE_GRAY
    ln.line.width = Pt(1.0)


def notes(slide, lines):
    tf = slide.notes_slide.notes_text_frame
    tf.text = "\n".join(lines)


def save(prs, outdir, name):
    path = outdir / name
    prs.save(path)
    print(f"  {path.name}")
    return path


# ---------------------------------------------------------------------------
# QC7つ道具（作図するもの）
# ---------------------------------------------------------------------------

def t_fishbone(outdir):
    """特性要因図。4M の大骨に中骨・小骨を付ける。"""
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "特性要因図（要因を網羅的に出す）", "要因を4Mで分解し、検証する順に並べ替える")

    spine_y = Inches(4.1)
    head_x = Inches(11.0)
    # 背骨
    arrow(s, Inches(1.0), spine_y, head_x, spine_y, color=NAVY, width=2.0)
    # 特性（結果）
    box(s, head_x, Inches(3.65), Inches(1.9), Inches(0.9), "特性\n（不良・ばらつき）",
        fill=NAVY, line=NAVY, color=WHITE, bold=True, size=11)

    # 大骨（4M）。上2つ・下2つ
    BONE_RUN = 1.5   # 大骨の水平方向の伸び（インチ）
    bones = [
        ("人（Man）", 2.4, True), ("機械（Machine）", 6.6, True),
        ("材料（Material）", 2.4, False), ("方法（Method）", 6.6, False),
    ]
    for label, bx, upper in bones:
        y_start = 1.95 if upper else 6.25
        y_end = 4.1  # 背骨のy
        arrow(s, Inches(bx), Inches(y_start), Inches(bx + BONE_RUN), Inches(y_end),
              color=BLUE, width=1.75)
        ly = Inches(1.52) if upper else Inches(6.4)
        box(s, Inches(bx - 0.7), ly, Inches(1.7), Inches(0.38), label,
            fill=LIGHT, line=BLUE, size=10, bold=True)

        # 中骨2本。始点は必ず大骨の線上に置く。
        # 大骨は (bx, y_start) → (bx+BONE_RUN, y_end) の直線なので、
        # 中骨のy から進行割合を逆算して x を求める。ここを固定値にすると骨を突き抜ける。
        for k in (0, 1):
            sy = (2.55 + 0.68 * k) if upper else (5.65 - 0.68 * k)
            t = (sy - y_start) / (y_end - y_start)
            sx = bx + BONE_RUN * t
            line(s, Inches(sx), Inches(sy), Inches(sx + 1.25), Inches(sy),
                 color=GRAY, width=1.0)
            # ラベルは大骨から離れる側に置く。上の骨は線の上、下の骨は線の下。
            # どちらも線の上に置くと、下側で大骨と文字が重なる
            label_y = (sy - 0.24) if upper else (sy + 0.05)
            textbox(s, Inches(sx + 0.1), Inches(label_y), Inches(1.25), Inches(0.2),
                    "中骨を記入", size=8, color=GRAY)

    notes(s, [
        "【特性要因図の使い方】",
        "・特性（結果）は1つに絞る。複数の不良をまとめて1枚に描かない",
        "・大骨は4M（人・機械・材料・方法）が基本。必要に応じて測定・環境を足して5M+1Eにする",
        "・中骨・小骨は「なぜそうなるか」で1〜2段掘る。1段で止めると対策が抽象的になる",
        "・末端が精神論（注意する・気をつける）になったら掘り足りない",
        "・出し切ったら、検証できる順に番号を振る。全部を検証しようとしない",
        "・現場の人を必ず入れて描く。1人で描くと現場感のない要因しか出ない",
        "",
        "【編集方法】中骨の線をコピーして増やす。テキストボックスは線の上に置く",
    ])
    return save(prs, outdir, "07_特性要因図.pptx")


# ---------------------------------------------------------------------------
# 新QC7つ道具
# ---------------------------------------------------------------------------

def t_affinity(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "親和図法（バラバラの意見を構造にする）",
          "現場の言葉をカードに出し、似たものを集めて全体像を作る")

    groups = ["グループ見出し1", "グループ見出し2", "グループ見出し3", "グループ見出し4"]
    cards = 4
    gw = Inches(2.95)
    gh = Inches(0.58 + 0.66 * cards + 0.14)
    for i, g in enumerate(groups):
        gx = Inches(0.5) + (i % 4) * Inches(3.15)
        gy = Inches(1.55)
        box(s, gx, gy, gw, gh, "", fill=PALE, line=BLUE)
        box(s, gx, gy, gw, Inches(0.42), g, fill=BLUE, line=BLUE,
            color=WHITE, bold=True, size=11)
        for k in range(cards):
            box(s, gx + Inches(0.14), gy + Inches(0.58 + 0.66 * k),
                gw - Inches(0.28), Inches(0.54), f"カード{k + 1}",
                fill=WHITE, line=LINE_GRAY, size=9, color=INK)

    sy = Inches(1.55) + gh + Inches(0.45)
    textbox(s, Inches(0.5), sy - Inches(0.32), Inches(12.33), Inches(0.26),
            "総括文（全体を1文で言い切る）", size=10, bold=True, color=NAVY)
    box(s, Inches(0.5), sy, Inches(12.33), Inches(0.72), "",
        fill=WHITE, line=NAVY)

    notes(s, [
        "【親和図法の使い方】",
        "・カードから始める。分類軸を先に決めると、結局いつもの分類になる",
        "・カードは1枚1事実。「〜が多い」ではなく、誰がいつ何を見たかを書く",
        "・グループの見出しは、中のカードの言葉を使って作る。抽象語で括らない",
        "・報告した本人を必ず入れる。伝聞だけで作ると文脈が抜ける",
        "・グループは4〜6。多すぎるなら括り直す",
        "・最後に全体を1文で言い切る。言えないならまだ構造になっていない",
        "",
        "【編集方法】カードはコピーして増やす。グループ枠の高さも一緒に伸ばす",
    ])
    return save(prs, outdir, "01_親和図法.pptx")


def t_relations(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "連関図法（絡み合う要因から真因を探す）",
          "要因どうしを矢印で結び、出る矢印の多い要因を真因として選ぶ")

    hub = box(s, Inches(5.65), Inches(3.6), Inches(2.4), Inches(1.05),
              "問題\n（中心に置く）", fill=NAVY, line=NAVY, color=WHITE,
              bold=True, size=12)

    pos = [
        (1.0, 1.7), (5.75, 1.62), (10.4, 1.7), (0.85, 3.72),
        (10.55, 3.72), (1.0, 5.72), (5.75, 5.95), (10.4, 5.72),
    ]
    w, h = 1.9, 0.72
    for i, (x, y) in enumerate(pos, start=1):
        f = box(s, Inches(x), Inches(y), Inches(w), Inches(h), f"要因{i}",
                fill=LIGHT, line=BLUE, size=10)
        # 始点・終点をそれぞれの箱の縁に取る。中心どうしを結ぶと線が箱の中に入り、
        # 中心の文字に矢尻が重なって読めなくなる
        x1, y1 = edge_point(rect_of(f), rect_of(hub))
        x2, y2 = edge_point(rect_of(hub), rect_of(f))
        arrow(s, x1, y1, x2, y2, color=GRAY, width=1.0)

    textbox(s, Inches(0.5), Inches(6.95), Inches(12.33), Inches(0.3),
            "矢印は「原因 → 結果」の向きで引く／出ている矢印が多い要因が真因の候補",
            size=9, color=GRAY)

    notes(s, [
        "【連関図法の使い方】",
        "・矢印は必ず「原因 → 結果」の向き。向きを間違えると症状に対策してしまう",
        "・出ている矢印が多い＝多くを引き起こしている＝真因の候補",
        "・入ってくる矢印が多い＝結果側。ここを叩いても再発する",
        "・要因を絞る。全部つなぐと矢印だらけで読めない。主要なものに限定する",
        "・特性要因図が「網羅」なら、連関図は「関係」。目的が違うので併用する",
        "",
        "【編集方法】不要な矢印は削除する。要因ボックスは複製して配置し直す",
    ])
    return save(prs, outdir, "02_連関図法.pptx")


def t_tree(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "系統図法（目的から手段へ展開する）",
          "目的を手段に分解し、末端が実行できる粒度になるまで掘る")

    x0 = Inches(0.5)
    goal = rbox(s, x0, Inches(3.5), Inches(2.2), Inches(0.9), "目的",
                fill=NAVY, line=NAVY, color=WHITE, bold=True, size=12)
    x1 = x0 + Inches(2.7)
    x2 = x1 + Inches(3.0)
    x3 = x2 + Inches(3.3)

    ys1 = [Inches(1.9), Inches(3.6), Inches(5.3)]
    for i, y in enumerate(ys1, start=1):
        box(s, x1, y, Inches(2.5), Inches(0.75), f"1次手段{i}", fill=LIGHT,
            line=BLUE, size=11, bold=True)
        arrow(s, x0 + Inches(2.2), Inches(3.95), x1, y + Inches(0.37),
              color=BLUE, width=1.25)
        for k in range(2):
            yy = y - Inches(0.32) + k * Inches(0.82)
            box(s, x2, yy, Inches(2.8), Inches(0.68), f"2次手段{i}-{k + 1}",
                fill=PALE, line=LINE_GRAY, size=10)
            arrow(s, x1 + Inches(2.5), y + Inches(0.37), x2, yy + Inches(0.34),
                  color=GRAY, width=1.0)
            box(s, x3, yy, Inches(2.9), Inches(0.68), "具体策（誰が・いつ・何を）",
                fill=WHITE, line=LINE_GRAY, size=9, color=GRAY)
            arrow(s, x2 + Inches(2.8), yy + Inches(0.34), x3, yy + Inches(0.34),
                  color=GRAY, width=1.0)

    notes(s, [
        "【系統図法の使い方】",
        "・左が目的、右が手段。右へ行くほど具体になる",
        "・各段で「そのために何をするか」を問う。答えが精神論なら展開が浅い",
        "・末端は「誰が・いつ・何を」まで書ける粒度にする。書けないならもう1段掘る",
        "・対策が全部『教育』になったら5E（排除・代替・隔離・工学・防護）の観点を足す",
        "・展開しきったら、マトリックス図法で効果と実現性を評価して絞る",
        "",
        "【編集方法】枝はグループごとに複製する。段数は3段で足りなければ4段に増やす",
    ])
    return save(prs, outdir, "03_系統図法.pptx")


def t_matrix(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "マトリックス図法（複数の案を軸で比べる）",
          "案を評価軸で採点し、選んだ理由を残す")

    cols = ["評価軸1\n効果", "評価軸2\n実現性", "評価軸3\nコスト", "評価軸4\n期間", "総合"]
    rows = ["対策案A", "対策案B", "対策案C", "対策案D"]
    x0, y0 = Inches(0.8), Inches(1.85)
    cw0, cw, rh = Inches(2.6), Inches(1.85), Inches(0.72)

    box(s, x0, y0, cw0, rh, "", fill=WHITE, line=LINE_GRAY)
    for j, c in enumerate(cols):
        box(s, x0 + cw0 + j * cw, y0, cw, rh, c, fill=NAVY, line=NAVY,
            color=WHITE, bold=True, size=10)
    for i, rname in enumerate(rows):
        yy = y0 + (i + 1) * rh
        box(s, x0, yy, cw0, rh, rname, fill=LIGHT, line=BLUE, size=11, bold=True)
        for j in range(len(cols)):
            box(s, x0 + cw0 + j * cw, yy, cw, rh, "", fill=WHITE,
                line=LINE_GRAY, size=13)

    ly = y0 + 5 * rh + Inches(0.35)
    textbox(s, x0, ly, Inches(11.5), Inches(0.28),
            "記号の意味：  ◎ = 5点（明確に優れる）   ○ = 3点（条件付きで可）   "
            "△ = 1点（懸念あり）   × = 0点（不可）", size=10, color=INK)
    textbox(s, x0, ly + Inches(0.36), Inches(11.5), Inches(0.28),
            "選定理由（1文で言い切る）：", size=10, bold=True, color=NAVY)
    box(s, x0, ly + Inches(0.66), Inches(11.5), Inches(0.55), "",
        fill=WHITE, line=LINE_GRAY)

    notes(s, [
        "【マトリックス図法の使い方】",
        "・全部◎になるのは軸が甘い証拠。相対評価にして差を付ける",
        "・軸に「実現性」を必ず入れる。効果だけで選ぶと実行されない案が残る",
        "・軸に重みを付ける場合は、重みも表に書く。頭の中で加重しない",
        "・記号だけ残して理由を残さないと、後から覆される。選定理由を必ず1文書く",
        "・落とした案も表に残す。検討したことが証拠になる",
        "",
        "【編集方法】行・列はセルを複製して増やす。総合列は手計算せず合計を記入する",
    ])
    return save(prs, outdir, "04_マトリックス図法.pptx")


def t_arrow(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "アローダイアグラム法（日程と順序を組む）",
          "作業の前後関係を並べ、最も長い経路を工期として押さえる")

    y_main = Inches(2.5)
    nodes_x = [Inches(0.9), Inches(3.3), Inches(5.7), Inches(8.1), Inches(10.5), Inches(12.2)]
    for i, x in enumerate(nodes_x):
        shape(s, MSO_SHAPE.OVAL, x, y_main, Inches(0.62), Inches(0.62), str(i + 1),
              fill=NAVY, line=NAVY, color=WHITE, bold=True, size=12)
    for i in range(len(nodes_x) - 1):
        x1 = nodes_x[i] + Inches(0.62)
        x2 = nodes_x[i + 1]
        arrow(s, x1, y_main + Inches(0.31), x2, y_main + Inches(0.31),
              color=ACCENT, width=2.0)
        mid = x1 + (x2 - x1) / 2 - Inches(0.6)
        textbox(s, mid, y_main - Inches(0.42), Inches(1.5), Inches(0.24),
                f"作業{chr(65 + i)}", size=10, bold=True, color=INK,
                align=PP_ALIGN.CENTER)
        textbox(s, mid, y_main + Inches(0.72), Inches(1.5), Inches(0.24),
                "○日", size=9, color=GRAY, align=PP_ALIGN.CENTER)

    # 並行作業（クリティカルパス外）。結合点番号は本線と重複させない
    y_sub = Inches(4.6)
    shape(s, MSO_SHAPE.OVAL, nodes_x[2], y_sub, Inches(0.62), Inches(0.62), "7",
          fill=BLUE, line=BLUE, color=WHITE, bold=True, size=12)
    arrow(s, nodes_x[1] + Inches(0.4), y_main + Inches(0.62), nodes_x[2],
          y_sub + Inches(0.31), color=GRAY, width=1.25)
    arrow(s, nodes_x[2] + Inches(0.62), y_sub + Inches(0.31),
          nodes_x[3] + Inches(0.2), y_main + Inches(0.62), color=GRAY, width=1.25)
    textbox(s, nodes_x[2] - Inches(0.3), y_sub + Inches(0.7), Inches(1.5), Inches(0.24),
            "並行作業", size=9, color=GRAY, align=PP_ALIGN.CENTER)

    textbox(s, Inches(0.9), Inches(6.25), Inches(11.5), Inches(0.28),
            "赤の経路がクリティカルパス／ここが1日延びると全体が1日延びる",
            size=10, bold=True, color=ACCENT)
    textbox(s, Inches(0.9), Inches(6.61), Inches(11.5), Inches(0.28),
            "灰の経路には余裕（フロート）がある／遅れても全体には響かない",
            size=10, color=GRAY)

    notes(s, [
        "【アローダイアグラム法の使い方】",
        "・矢印が作業、丸が結合点（作業の区切り）。矢印の上に作業名、下に所要日数を書く",
        "・最も長い経路がクリティカルパス。ここを短縮しないと全体は縮まない",
        "・余裕のある経路に人を張り付けても工期は縮まない。まずクリティカルパスを見る",
        "・並行できる作業を洗い出してから並べる。順番に並べるだけでは短縮できない",
        "・ガントチャートは進捗管理向き、アローダイアグラムは順序の設計向き。役割が違う",
        "",
        "【編集方法】結合点と矢印を複製して増やす。クリティカルパスだけ赤・太線にする",
    ])
    return save(prs, outdir, "05_アローダイアグラム法.pptx")


def t_pdpc(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "PDPC法（失敗した場合の迂回路を描く）",
          "うまくいかない場合を先に想定し、代替手段を決めておく")

    x_main = Inches(1.5)
    w, h = Inches(2.9), Inches(0.78)
    steps = ["出発点（現状）", "手段1を実行", "手段2を実行", "目標（達成状態）"]
    ys = [Inches(1.7), Inches(3.0), Inches(4.3), Inches(5.6)]
    for i, (st, y) in enumerate(zip(steps, ys)):
        last = i == len(steps) - 1
        rbox(s, x_main, y, w, h, st,
             fill=NAVY if (i == 0 or last) else LIGHT,
             line=NAVY if (i == 0 or last) else BLUE,
             color=WHITE if (i == 0 or last) else INK, bold=True, size=11)
        if i < len(steps) - 1:
            arrow(s, x_main + w / 2, y + h, x_main + w / 2, ys[i + 1],
                  color=NAVY, width=1.5)

    # 分岐と迂回路
    x_alt = Inches(6.3)
    x_alt2 = Inches(9.8)
    alt_w = Inches(2.9)
    for i, y in enumerate(ys[1:3], start=1):
        shape(s, MSO_SHAPE.DIAMOND, x_alt, y - Inches(0.07), Inches(2.3), Inches(0.92),
              "うまくいくか", fill=WHITE, line=ACCENT, size=9, color=INK)
        arrow(s, x_main + w, y + h / 2, x_alt, y + Inches(0.39),
              color=GRAY, width=1.25)
        # 「想定外」は矢印の上に置く。ボックスに重ねると読めない
        textbox(s, x_main + w + Inches(0.35), y - Inches(0.02), Inches(1.2), Inches(0.22),
                "想定外", size=8, bold=True, color=ACCENT)
        alt = rbox(s, x_alt2, y, alt_w, h, f"代替手段{i}", fill=PALE,
                   line=ACCENT, size=11)
        arrow(s, x_alt + Inches(2.3), y + Inches(0.39), x_alt2, y + h / 2,
              color=ACCENT, width=1.25)
        # 代替手段から本線の次のステップへ戻す。
        # 縦に降ろすと下段の代替手段ボックスを貫通するので、左下へ1本で結ぶ。
        # 着地点を次のステップの上寄り（+0.2）にすると、途中のひし形の上を通り抜ける
        arrow(s, x_alt2, y + h, x_main + w, ys[i + 1] + Inches(0.2),
              color=ACCENT, width=1.25)

    notes(s, [
        "【PDPC法の使い方】",
        "・PDPC = Process Decision Program Chart。「思い通りに進まない場合」を先に描く",
        "・本線（うまくいく道）を先に描き、後から分岐と迂回路を足す",
        "・分岐は「うまくいくか」ではなく「何が起きたら止まるか」で考えると具体になる",
        "・代替手段は必ず本線に戻す。戻り先が無い代替手段は使えない",
        "・分岐は多くても3箇所。全ての失敗を描こうとすると読めなくなる",
        "・系統図法で手段を出し、PDPCでその手段の失敗に備える、という順で使う",
        "",
        "【編集方法】分岐のひし形と代替手段は1組をコピーして増やす",
    ])
    return save(prs, outdir, "06_PDPC法.pptx")


def t_matrix_data(outdir):
    prs = new_deck()
    s = blank_slide(prs)
    frame(s, "マトリックスデータ解析法（多変量を少数の軸に要約する）",
          "多くの評価項目を2軸に圧縮し、位置づけの違いを見る")

    # 象限
    ox, oy = Inches(2.4), Inches(1.75)
    ow, oh = Inches(8.4), Inches(4.9)
    box(s, ox, oy, ow, oh, "", fill=WHITE, line=LINE_GRAY)
    line(s, ox, oy + oh / 2, ox + ow, oy + oh / 2, color=GRAY, width=1.0, dash=True)
    line(s, ox + ow / 2, oy, ox + ow / 2, oy + oh, color=GRAY, width=1.0, dash=True)

    arrow(s, ox, oy + oh, ox, oy - Inches(0.15), color=NAVY, width=1.5)
    arrow(s, ox, oy + oh, ox + ow + Inches(0.15), oy + oh, color=NAVY, width=1.5)
    textbox(s, ox - Inches(1.85), oy + oh / 2 - Inches(0.15), Inches(1.7), Inches(0.3),
            "第2主成分（軸2）", size=10, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)
    textbox(s, ox + ow / 2 - Inches(1.0), oy + oh + Inches(0.18), Inches(2.0), Inches(0.3),
            "第1主成分（軸1）", size=10, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    for label, px, py, color in [
        ("対象A", Inches(3.6), Inches(2.6), NAVY),
        ("対象B", Inches(8.6), Inches(2.9), NAVY),
        ("対象C", Inches(4.3), Inches(5.3), BLUE),
        ("対象D", Inches(9.4), Inches(5.0), BLUE),
        ("自社", Inches(6.3), Inches(3.9), ACCENT),
    ]:
        shape(s, MSO_SHAPE.OVAL, px, py, Inches(0.22), Inches(0.22), "",
              fill=color, line=color)
        textbox(s, px + Inches(0.28), py - Inches(0.06), Inches(1.2), Inches(0.24),
                label, size=10, bold=True, color=color)

    for q, qx, qy in [("第2象限", ox + Inches(0.15), oy + Inches(0.12)),
                      ("第1象限", ox + ow - Inches(1.1), oy + Inches(0.12)),
                      ("第3象限", ox + Inches(0.15), oy + oh - Inches(0.38)),
                      ("第4象限", ox + ow - Inches(1.1), oy + oh - Inches(0.38))]:
        textbox(s, qx, qy, Inches(1.0), Inches(0.24), q, size=9, color=GRAY)

    notes(s, [
        "【マトリックスデータ解析法の使い方】",
        "・新QC7つ道具で唯一、数値データを使う道具。主成分分析にあたる",
        "・多数の評価項目を、情報量の多い2軸に圧縮して散布図にする",
        "・軸には必ず意味を読み取って名前を付ける。「第1主成分」のままでは説明できない",
        "・軸の意味は寄与率と因子負荷量から読む。感覚で名付けない",
        "・寄与率の合計が低い（目安7割未満）なら2軸で語れない。軸を増やすか手法を変える",
        "・計算は Excel の分析ツールか統計ソフトで行う。この図は結果を貼る枠",
        "",
        "【編集方法】点はコピーして増やす。軸名は必ず実データから読み取って書き換える",
    ])
    return save(prs, outdir, "07_マトリックスデータ解析法.pptx")


QC7_BUILDERS = [t_fishbone]
N7_BUILDERS = [t_affinity, t_relations, t_tree, t_matrix, t_arrow, t_pdpc,
               t_matrix_data]


def _assets_templates():
    """図テンプレートの置き場所。**このスキルの中。**

    スキルは `~/.claude/skills/<name>` から参照される。
    **その外に置いたものは、スキルの一部ではない。**
    実際に、リポジトリ直下に置いたところ
    「スキルだけを配ると雛形が0件」「規約の相対パスが空振り」になった。

    `__file__` は登録のリンク越しに呼ばれるので、**resolve() してから辿る。**
    """
    return pathlib.Path(__file__).resolve().parent.parent / "assets" / "templates"


def main():
    import argparse

    ap = argparse.ArgumentParser(description="作図系テンプレート（PowerPoint）を生成する")
    ap.add_argument("qc7dir", nargs="?", default=None)
    ap.add_argument("n7dir", nargs="?", default=None)
    ap.add_argument("--font", default=FONT,
                    help="日本語フォント名。**資料テンプレートに合わせる。**"
                         "read_template_style.py で調べた値を渡す")
    ap.add_argument("--latin", default=None,
                    help="欧文フォント名。省くと日本語と同じものを使う")
    args = ap.parse_args()

    # フォントは本スキルで固定しない。資料テンプレートが決める
    globals()["FONT"] = args.font
    globals()["LATIN"] = args.latin

    base = _assets_templates()
    qc7 = Path(args.qc7dir) if args.qc7dir else base / "qc7"
    n7 = Path(args.n7dir) if args.n7dir else base / "n7"
    qc7.mkdir(parents=True, exist_ok=True)
    n7.mkdir(parents=True, exist_ok=True)

    print(f"フォント: 日本語={FONT} / 欧文={LATIN or FONT}")
    print(f"QC7つ道具テンプレート（PowerPoint）を生成: {qc7}")
    for fn in QC7_BUILDERS:
        fn(qc7)
    print(f"新QC7つ道具テンプレート（PowerPoint）を生成: {n7}")
    for fn in N7_BUILDERS:
        fn(n7)
    print(f"完了: {len(QC7_BUILDERS) + len(N7_BUILDERS)} 件")


if __name__ == "__main__":
    main()
