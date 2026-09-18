#!/usr/bin/env python3
"""出てきた pptx の中身を、**1コマンドで**要約する。

    python3 scripts/probe_pptx.py deck.pptx
    python3 scripts/probe_pptx.py deck.pptx --slide 7     # 1枚だけ詳しく

## なぜあるか

変換を直したあと「本当に直ったか」を確かめるのに、
**毎回その場で unzip して XML を grep していた。**
時間もかかるし、確かめるたびに違う書き方になって**比べられない。**

**python-pptx で全部読める。**`<a:ea>`（和文書体）だけは
オブジェクトに出ていないので `_rPr` から読む。**unzip は要らない。**

## 何を見るか

| 見るもの | 壊れていると |
|---|---|
| **枠の自動拡張** | 折り返した枠が下へ伸び、隣・下の要素と重なる |
| **折り返しの有無** | `nowrap` が効かず、枠外へこぼれる |
| **文字サイズ** | 既定（18pt）が混じると、体裁を読めていない |
| **書体（欧文／和文）** | 和文が欧文書体に渡ると字幅が変わる。**フォント欄の表示も食い違う** |
| **表のセル余白・縦位置** | 既定のままだと、収まる文字が折り返して行が伸びる |

**数えるだけで、直しはしない。**直すのは HTML（`SKILL.md` の絶対の制約）。
"""

import argparse
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.oxml.ns import qn

import checked
from html_abs import _JA_RE

# PowerPoint の既定。**これが出たら、こちらが指定できていない**
DEFAULT_PT = 18.0
DEFAULT_MARGIN = (91440, 91440)     # 左右 0.1in


def runs_of(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            yield from p.runs
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    yield from p.runs


def ea_of(run):
    el = run.font._rPr.find(qn("a:ea"))
    return el.get("typeface") if el is not None else None


def main():
    ap = argparse.ArgumentParser(description="pptx の中身を要約する（読むだけ）。絵に違和感があるのに原因が分からないときに使う")
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int, help="この枚だけ詳しく出す")
    args = ap.parse_args()

    src = Path(args.pptx)
    if not src.exists():
        sys.exit(f"見つからない: {src}")
    prs = Presentation(str(src))

    autofit, wrap, sizes = Counter(), Counter(), Counter()
    faces, margins, anchors = Counter(), Counter(), Counter()
    findings, n_shapes, n_runs = [], 0, 0

    for no, slide in enumerate(prs.slides, 1):
        if args.slide and no != args.slide:
            continue
        for sh in slide.shapes:
            n_shapes += 1
            if sh.has_text_frame:
                a = str(sh.text_frame.auto_size)
                autofit[a] += 1
                wrap[str(sh.text_frame.word_wrap)] += 1
                if "NONE" not in a:
                    findings.append((no, "枠の自動拡張",
                                     f"{sh.shape_id} が {a}。枠が伸びる"))
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        m = (cell.margin_left, cell.margin_right)
                        margins[m] += 1
                        anchors[str(cell.vertical_anchor)] += 1
                        if m == DEFAULT_MARGIN:
                            findings.append((no, "表セルの余白",
                                             "既定 0.1in のまま。CSS を読めていない"))
            for r in runs_of(sh):
                n_runs += 1
                pt = r.font.size.pt if r.font.size else None
                sizes[pt] += 1
                if pt == DEFAULT_PT:
                    findings.append((no, "既定の文字サイズ",
                                     f"18pt: {r.text[:20]}"))
                latin, ea = r.font.name, ea_of(r)
                faces[(latin, ea)] += 1
                if _JA_RE.search(r.text or ""):
                    if not ea:
                        findings.append((no, "和文の書体が無い", r.text[:20]))
                    elif latin and latin != ea and not any(
                            c.isascii() and c.isalnum() for c in r.text):
                        # 和文だけの run。latin にも和文書体が入っていないと、
                        # **描画は正しくてもフォント欄には欧文書体と出る**
                        findings.append((no, "フォント欄が食い違う",
                                         f"{r.text[:16]} latin={latin} ea={ea}"))

    def line(label, counter, top=6):
        items = counter.most_common(top)
        body = " / ".join(f"{k}×{v}" for k, v in items) or "—"
        extra = "" if len(counter) <= top else f" ほか{len(counter) - top}種"
        print(f"| {label} | {body}{extra} |")

    where = f"（{args.slide}枚目のみ）" if args.slide else ""
    print(f"pptx の中身: {len(prs.slides)} 枚  図形 {n_shapes}  run {n_runs} {where}\n")
    print("| 見るもの | 値 |")
    print("|---|---|")
    line("枠の自動拡張", autofit)
    line("折り返し", wrap)
    line("文字サイズ(pt)", sizes)
    line("書体 (欧文, 和文)", faces)
    line("表セル余白", margins)
    line("表セル縦位置", anchors)

    if findings:
        print(f"\n## 気になるところ: {len(findings)} 件\n")
        print("| 枚 | 種別 | 内容 |")
        print("|---|---|---|")
        seen = set()
        for no, kind, msg in findings:
            key = (no, kind)
            if key in seen:
                continue
            seen.add(key)
            n = sum(1 for f in findings if (f[0], f[1]) == key)
            print(f"| {no} | {kind} | {msg}{'' if n == 1 else f'（ほか{n - 1}件）'} |")
        print("\n**直すのは HTML。**出てきた pptx を手で直さない。")

    print("\n**これは中身の要約であって、見た目の検査ではない。**")
    print("重なり・溢れは `fit_check.py`、絵は `qa_render.py` で見る。")

    return checked.summary("probe_pptx", n_runs, "run", len(findings), {
        "枚": len(prs.slides),
        "図形": n_shapes,
        "元": src.name,
    })


if __name__ == "__main__":
    sys.exit(main())
