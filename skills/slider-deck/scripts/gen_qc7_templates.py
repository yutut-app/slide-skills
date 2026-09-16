#!/usr/bin/env python3
"""QC7つ道具のうち、数値を扱うもののテンプレート（Excel）を生成する。

1テンプレ = 1ファイル = 1シート。
道具の名前と用途は mfg-improvement-frameworks スキルの references/30_qc7.md に揃える。
用語がずれると2つのスキルが噛み合わなくなるため、勝手に言い換えない。

特性要因図は作図が主なので PowerPoint 側（gen_qc7_pptx.py）で作る。
グラフは charts/ 配下のテンプレートを使う。

使い方:
    python3 scripts/gen_qc7_templates.py [出力先ディレクトリ]
    既定の出力先は assets/templates/qc7
"""

import re
import sys
import pathlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.trendline import Trendline
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F4E79"
BLUE = "4472C4"
LIGHT = "BDD7EE"
GRAY = "808080"
ACCENT = "C00000"
INK = "333333"
FONT = "游ゴシック"   # 日本語
LATIN = None          # 欧文。None なら日本語と同じ

HEAD_FILL = PatternFill("solid", fgColor=NAVY)
NOTE_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 入力欄。記入する場所を示す
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, title, purpose, min_n=None):
    ws.sheet_view.showGridLines = False
    # 印刷・PDF化で表とグラフが1ページに収まるようにする。
    # 既定の縦向きだとグラフだけ次ページに落ちる
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws["A2"] = purpose
    ws["A2"].font = Font(name=FONT, size=10, color=INK)
    if min_n:
        ws["A3"] = f"必要件数： {min_n}"
        ws["A3"].font = Font(name=FONT, size=10, bold=True, color=ACCENT)
    ws.column_dimensions["A"].width = 3


def table(ws, rows, first_row=5, first_col=2, input_cols=()):
    ncols = max(len(r) for r in rows)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            ws.cell(row=first_row + i, column=first_col + j, value=val)
    for c in range(first_col, first_col + ncols):
        cell = ws.cell(row=first_row, column=c)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    last = first_row + len(rows) - 1
    for r in range(first_row + 1, last + 1):
        for c in range(first_col, first_col + ncols):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.border = BORDER
            if c > first_col:
                cell.alignment = Alignment(horizontal="right")
            if (c - first_col) in input_cols:
                cell.fill = INPUT_FILL
    for c in range(first_col, first_col + max(ncols, 4)):
        ws.column_dimensions[get_column_letter(c)].width = 15
    return {"first_row": first_row, "first_col": first_col, "last_row": last, "ncols": ncols}


def notes(ws, ref, lines, col_offset=0):
    r = ref["last_row"] + 2
    c = ref["first_col"] + col_offset
    ws.cell(row=r, column=c, value="使い方と注意").font = Font(
        name=FONT, size=10, bold=True, color=NAVY
    )
    for i, line in enumerate(lines, start=1):
        cell = ws.cell(row=r + i, column=c, value=line)
        cell.font = Font(name=FONT, size=9, color=INK)
        cell.fill = NOTE_FILL



def chart_font(ch):
    """後方互換のための空処理。

    openpyxl の ChartBase には txPr が無く（ChartSpace 側の要素）、代入しても
    黙って無視される。書体は保存後に inject_chart_font() で入れる。
    """
    return ch


def inject_chart_font(path):
    """保存した xlsx のグラフに、既定の書体を入れる。

    グラフの中の文字（タイトル・軸・凡例・ラベル）は、指定しないとテーマの
    既定になる。**表はテンプレートの書体、グラフだけ別**という状態を避ける。
    ea を書かないと日本語に効かない点はスライドと同じ。

    <c:txPr> は <c:chart> の後ろに置く必要がある（スキーマ上の順序）。
    """
    import shutil
    import zipfile

    latin = LATIN or FONT
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    # openpyxl は接頭辞なし（既定名前空間）で書く。PowerPoint 由来だと c: が付く。
    # **閉じタグの形を実測して合わせる。**決め打ちすると黙って何も起きない
    def make(prefix):
        p = f"{prefix}:" if prefix else ""
        return (
            f'<{p}txPr><a:bodyPr xmlns:a="{A}"/><a:lstStyle xmlns:a="{A}"/>'
            f'<a:p xmlns:a="{A}"><a:pPr><a:defRPr>'
            f'<a:latin typeface="{latin}"/><a:ea typeface="{FONT}"/>'
            f'<a:cs typeface="{latin}"/>'
            f'</a:defRPr></a:pPr><a:endParaRPr lang="ja-JP"/></a:p></{p}txPr>'
        )

    import shutil
    import zipfile

    src = zipfile.ZipFile(path)
    tmp = str(path) + ".tmp"
    touched = 0
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if re.match(r"xl/charts/chart\d+\.xml$", item.filename):
                x = data.decode("utf-8")
                for prefix, close in (("", "</chart>"), ("c", "</c:chart>")):
                    if close in x and "txPr>" not in x:
                        x = x.replace(close, close + make(prefix), 1)
                        data = x.encode("utf-8")
                        touched += 1
                        break
            out.writestr(item, data)
    src.close()
    shutil.move(tmp, path)
    return touched


def value_labels(ch, percent=False, cat_name=False):
    """データラベルを付ける。表示する項目を全部明示する。

    showVal だけ立てて他を未指定にすると、レンダラ既定で系列名やカテゴリ名まで
    出て「寸法不良（外径）: 件数: 186」のような読めないラベルになる。
    """
    dl = DataLabelList()
    dl.showVal = not percent
    dl.showPercent = percent
    dl.showCatName = cat_name
    dl.showSerName = False
    dl.showLegendKey = False
    dl.showBubbleSize = False
    ch.dLbls = dl
    return ch


def solid(series, hexcolor, line=False):
    if line:
        series.graphicalProperties.line.solidFill = hexcolor
        series.graphicalProperties.line.width = 22000
    else:
        series.graphicalProperties.solidFill = hexcolor
        series.graphicalProperties.line.noFill = True
    return series


def save(wb, outdir, name):
    path = outdir / name
    wb.save(path)
    inject_chart_font(path)
    print(f"  {path.name}")
    return path


# ---------------------------------------------------------------------------

def t_pareto(outdir):
    """パレート図。棒（件数の降順）＋折れ線（累積構成比）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "パレート図"
    header(ws, "パレート図（どこが大きいかを見る）",
           "不良や苦情を項目別に分け、影響の大きい項目を絞り込む。現状把握で使う。",
           "最低50件、望ましくは100件以上")

    items = [("寸法不良（外径）", 186), ("キズ・打痕", 152), ("組付け位置ずれ", 118),
             ("表面粗さ不良", 74), ("刻印かすれ", 41), ("その他", 29)]
    total = sum(v for _, v in items)
    rows = [["項目", "件数", "構成比", "累積構成比"]]
    cum = 0
    for name, v in items:
        cum += v
        rows.append([name, v, round(v / total * 100, 1), round(cum / total * 100, 1)])
    rows.append(["合計", total, 100.0, None])
    ref = table(ws, rows, input_cols=(0, 1))
    ws.column_dimensions["B"].width = 22

    notes(ws, ref, [
        "・件数の降順に並べる。「その他」は件数に関わらず必ず最後に置く",
        "・分類が抽象的だと毎回同じ項目が1位になり対策が効かない。具体化する",
        "・上位2〜3項目で全体の70〜80%を占めるなら、そこに絞って対策する",
        "・前後比較するときは期間の長さと生産数をそろえる。揃えないと改善が見えない",
        "・累積構成比の折れ線は第2軸（0〜100%）に置く",
    ])

    data_end = ref["last_row"] - 1  # 合計行はグラフに含めない
    bar = BarChart()
    bar.type = "col"
    bar.gapWidth = 20
    bar.add_data(Reference(ws, min_col=3, min_row=ref["first_row"], max_row=data_end),
                 titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=ref["first_row"] + 1, max_row=data_end))
    solid(bar.series[0], NAVY)
    value_labels(bar)
    bar.y_axis.title = "件数"

    line = LineChart()
    line.add_data(Reference(ws, min_col=5, min_row=ref["first_row"], max_row=data_end),
                  titles_from_data=True)
    solid(line.series[0], ACCENT, line=True)
    line.series[0].smooth = False
    line.series[0].marker = Marker(symbol="circle", size=6)
    line.y_axis.axId = 200
    line.y_axis.title = "累積構成比(%)"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 100
    line.y_axis.crosses = "max"
    bar += line
    bar.title = "不良項目のパレート図"
    bar.width, bar.height = 22, 12
    chart_font(bar)
    ws.add_chart(bar, "G5")
    return save(wb, outdir, "01_パレート図.xlsx")


def t_histogram(outdir):
    """ヒストグラム。度数表＋棒（間隔0）＋規格線の記入欄。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "ヒストグラム"
    header(ws, "ヒストグラム（ばらつきの形を見る）",
           "測定値の分布の形を見て、工程の状態と規格に対する余裕を判断する。",
           "最低30件、望ましくは100件以上")

    bins = [("9.90〜9.94", 2), ("9.94〜9.98", 7), ("9.98〜10.02", 19),
            ("10.02〜10.06", 31), ("10.06〜10.10", 24), ("10.10〜10.14", 11),
            ("10.14〜10.18", 4), ("10.18〜10.22", 2)]
    rows = [["区間", "度数"]] + [[a, b] for a, b in bins]
    ref = table(ws, rows, input_cols=(0, 1))
    ws.column_dimensions["B"].width = 16

    # 規格値の記入欄
    r = ref["last_row"] + 2
    ws.cell(row=r, column=2, value="規格下限(LSL)").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r, column=3, value=9.95).fill = INPUT_FILL
    ws.cell(row=r + 1, column=2, value="規格上限(USL)").font = Font(name=FONT, size=10, bold=True)
    ws.cell(row=r + 1, column=3, value=10.15).fill = INPUT_FILL
    for rr in (r, r + 1):
        ws.cell(row=rr, column=3).border = BORDER

    notes(ws, {"last_row": r + 1, "first_col": 2}, [
        "・区間の数は データ数の平方根 が目安。30件なら5〜6、100件なら10前後",
        "・区間の幅は全区間で同じにする。幅が違うと形が嘘になる",
        "・棒の間隔は0にする（連続量なので隙間を空けない）",
        "・山が2つ＝異なる条件が混ざっている。層別してから作り直す",
        "・片側に偏る＝選別や除去が行われている疑い",
        "・離れ小島＝異常値か別ロットの混入。生データに戻って確認する",
        "・規格線をグラフに引き、規格をはみ出していないか見る",
    ])

    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 0  # 連続量なので隙間を作らない
    ch.add_data(Reference(ws, min_col=3, min_row=ref["first_row"], max_row=ref["last_row"]),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=ref["first_row"] + 1,
                                max_row=ref["last_row"]))
    solid(ch.series[0], BLUE)
    ch.series[0].graphicalProperties.line.solidFill = "FFFFFF"
    ch.legend = None
    ch.title = "外径寸法の分布"
    ch.y_axis.title = "度数"
    ch.x_axis.title = "測定値(mm)"
    ch.width, ch.height = 22, 12
    chart_font(ch)
    ws.add_chart(ch, "F5")
    return save(wb, outdir, "02_ヒストグラム.xlsx")


def t_control_chart(outdir):
    """X̄-R管理図。X̄管理図とR管理図を縦に並べる。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "X-R管理図"
    header(ws, "X̅-R管理図（時間的な安定性を見る）",
           "工程が安定しているか、異常な変動が起きていないかを時系列で監視する。",
           "最低20〜25群、望ましくは25群以上（1群あたり4〜5個）")

    groups = [
        (10.02, 10.05, 9.98, 10.01), (10.00, 10.03, 10.06, 9.99),
        (9.97, 10.01, 10.04, 10.02), (10.05, 10.02, 9.99, 10.03),
        (10.01, 9.98, 10.02, 10.00), (10.03, 10.06, 10.01, 10.04),
        (9.99, 10.02, 10.00, 9.97), (10.04, 10.01, 10.03, 10.05),
        (10.00, 10.04, 10.02, 10.01), (10.02, 9.99, 10.01, 10.03),
        (10.06, 10.03, 10.05, 10.02), (9.98, 10.01, 9.99, 10.00),
        (10.03, 10.00, 10.02, 10.04), (10.01, 10.05, 10.03, 10.02),
        (10.00, 10.02, 9.98, 10.01), (10.04, 10.02, 10.06, 10.03),
        (10.02, 10.00, 10.01, 9.99), (10.05, 10.03, 10.02, 10.04),
        (9.99, 10.01, 10.00, 10.02), (10.03, 10.05, 10.04, 10.01),
        (10.01, 10.02, 10.03, 10.00), (10.02, 10.04, 10.01, 10.03),
        (10.00, 9.98, 10.02, 10.01), (10.04, 10.03, 10.05, 10.02),
        (10.02, 10.01, 10.00, 10.03),
    ]
    xbars = [round(sum(g) / len(g), 4) for g in groups]
    ranges = [round(max(g) - min(g), 4) for g in groups]
    gx = round(sum(xbars) / len(xbars), 4)   # X̄の平均（中心線）
    rbar = round(sum(ranges) / len(ranges), 4)
    # n=4 の管理図係数。JIS Z 9021 の表による
    A2, D3, D4 = 0.729, 0.0, 2.282
    ucl_x, lcl_x = round(gx + A2 * rbar, 4), round(gx - A2 * rbar, 4)
    ucl_r, lcl_r = round(D4 * rbar, 4), round(D3 * rbar, 4)

    rows = [["群No", "X̅", "UCL", "CL", "LCL", "R", "UCL(R)", "CL(R)", "LCL(R)"]]
    for i, (xb, rg) in enumerate(zip(xbars, ranges), start=1):
        rows.append([i, xb, ucl_x, gx, lcl_x, rg, ucl_r, rbar, lcl_r])
    ref = table(ws, rows, input_cols=(0, 1, 5))

    notes(ws, ref, [
        f"・管理限界は n=4 の係数（A2={A2}, D3={D3}, D4={D4}）で算出している。群の大きさを変えたら係数も変える",
        "・管理限界線は規格線ではない。工程自身のばらつきから計算する値であり、両者を混同しない",
        "・異常判定：限界線の外に出る／連続9点が中心線の同じ側／連続6点が上昇か下降／連続14点が交互に増減",
        "・Rが先に安定していないと X̅ の限界線は意味を持たない。R管理図から先に見る",
        "・群の取り方（時間・ロット・号機）で見えるものが変わる。何で群を作ったか必ず記録する",
    ])

    # X̄管理図
    ch1 = LineChart()
    ch1.add_data(Reference(ws, min_col=3, max_col=6, min_row=ref["first_row"],
                           max_row=ref["last_row"]), titles_from_data=True)
    ch1.set_categories(Reference(ws, min_col=2, min_row=ref["first_row"] + 1,
                                 max_row=ref["last_row"]))
    solid(ch1.series[0], NAVY, line=True)
    ch1.series[0].marker = Marker(symbol="circle", size=5)
    for idx, col in zip((1, 3), (ACCENT, ACCENT)):
        solid(ch1.series[idx], col, line=True)
        ch1.series[idx].graphicalProperties.line.dashStyle = "dash"
    solid(ch1.series[2], GRAY, line=True)
    for s in ch1.series:
        s.smooth = False
    ch1.title = "X̅管理図（群平均の推移）"
    ch1.y_axis.title = "X̅"
    ch1.width, ch1.height = 24, 9
    chart_font(ch1)
    ws.add_chart(ch1, "K5")

    # R管理図
    ch2 = LineChart()
    ch2.add_data(Reference(ws, min_col=7, max_col=10, min_row=ref["first_row"],
                           max_row=ref["last_row"]), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=ref["first_row"] + 1,
                                 max_row=ref["last_row"]))
    solid(ch2.series[0], BLUE, line=True)
    ch2.series[0].marker = Marker(symbol="circle", size=5)
    solid(ch2.series[1], ACCENT, line=True)
    ch2.series[1].graphicalProperties.line.dashStyle = "dash"
    solid(ch2.series[2], GRAY, line=True)
    solid(ch2.series[3], ACCENT, line=True)
    ch2.series[3].graphicalProperties.line.dashStyle = "dash"
    for s in ch2.series:
        s.smooth = False
    ch2.title = "R管理図（群内のばらつきの推移）"
    ch2.y_axis.title = "R"
    ch2.width, ch2.height = 24, 9
    chart_font(ch2)
    ws.add_chart(ch2, "K25")
    return save(wb, outdir, "03_X-R管理図.xlsx")


def t_scatter_stratified(outdir):
    """層別散布図。条件別に系列を分け、層別で相関が変わることを見る。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "層別散布図"
    header(ws, "散布図（2つの量の関係／層別つき）",
           "2つの量に関係があるかを見る。条件で層別し、見かけの相関を見破る。",
           "最低30組、望ましくは50組以上")

    a_x = [18.2, 19.1, 20.4, 21.0, 22.3, 23.1, 24.0, 25.2, 26.1, 27.4,
           28.0, 29.3, 30.1, 31.2, 32.0]
    a_y = [0.82, 0.88, 0.91, 0.95, 1.02, 1.05, 1.11, 1.18, 1.21, 1.29,
           1.31, 1.38, 1.42, 1.49, 1.52]
    b_x = [18.5, 19.4, 20.1, 21.3, 22.0, 23.4, 24.2, 25.0, 26.3, 27.1,
           28.4, 29.0, 30.3, 31.1, 32.4]
    b_y = [1.62, 1.58, 1.71, 1.65, 1.79, 1.74, 1.88, 1.82, 1.95, 1.90,
           2.03, 1.98, 2.11, 2.06, 2.19]
    rows = [["炉内温度(℃)", "反り量_号機A", "炉内温度(℃)", "反り量_号機B"]]
    for i in range(15):
        rows.append([a_x[i], a_y[i], b_x[i], b_y[i]])
    ref = table(ws, rows, input_cols=(0, 1, 2, 3))

    notes(ws, ref, [
        "・全体でひとまとまりに見えても、層別すると2つの群に分かれることがある",
        "・逆に、層別すると相関が消えることもある。層別前の相関を鵜呑みにしない",
        "・層別のキー（号機・ロット・作業者・時間帯）は記録時に取っておく。後から作れない",
        "・相関があっても因果とは限らない。第3の要因を疑う",
        "・近似直線と決定係数は目安。R²が高くても外挿はしない",
    ])

    ch = ScatterChart()
    ch.style = None
    for col_x, col_y, color in ((2, 3, NAVY), (4, 5, ACCENT)):
        xr = Reference(ws, min_col=col_x, min_row=ref["first_row"] + 1, max_row=ref["last_row"])
        yr = Reference(ws, min_col=col_y, min_row=ref["first_row"], max_row=ref["last_row"])
        s = Series(yr, xr, title_from_data=True)
        s.marker = Marker(symbol="circle", size=6)
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.noFill = True
        s.graphicalProperties.line.noFill = True
        s.trendline = Trendline(trendlineType="linear", dispRSqr=True)
        ch.series.append(s)
    ch.title = "炉内温度と反り量（号機で層別）"
    ch.x_axis.title = "炉内温度(℃)"
    ch.y_axis.title = "反り量(mm)"
    ch.width, ch.height = 22, 13
    chart_font(ch)
    ws.add_chart(ch, "G5")
    return save(wb, outdir, "04_層別散布図.xlsx")


def t_stratification(outdir):
    """層別集計表。4M の軸で分けて差を見る。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "層別"
    header(ws, "層別（条件で分けて差を見る）",
           "同じデータを条件で分け、どの条件に問題が集中しているかを見る。",
           "1水準あたり最低10件、望ましくは30件以上")

    rows = [["層別キー", "水準", "生産数", "不良数", "不良率(%)", "件数は十分か"]]
    data = [
        ("設備(Machine)", "1号機", 4200, 63, None, None),
        ("設備(Machine)", "2号機", 4100, 41, None, None),
        ("設備(Machine)", "3号機", 3900, 128, None, None),
        ("作業者(Man)", "Aさん", 4050, 58, None, None),
        ("作業者(Man)", "Bさん", 4080, 61, None, None),
        ("作業者(Man)", "Cさん", 4070, 113, None, None),
        ("材料ロット(Material)", "L-101", 6100, 92, None, None),
        ("材料ロット(Material)", "L-102", 6100, 140, None, None),
        ("時間帯(Method)", "昼勤", 8200, 98, None, None),
        ("時間帯(Method)", "夜勤", 4000, 134, None, None),
    ]
    for d in data:
        rows.append(list(d))
    ref = table(ws, rows, input_cols=(0, 1, 2, 3))
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["G"].width = 16

    # 不良率と件数判定を数式で入れる。手計算にしない
    for r in range(ref["first_row"] + 1, ref["last_row"] + 1):
        ws.cell(row=r, column=6, value=f"=IF(D{r}=0,\"\",ROUND(E{r}/D{r}*100,2))")
        ws.cell(row=r, column=7, value=f"=IF(D{r}>=30,\"OK\",\"件数不足\")")
        ws.cell(row=r, column=6).font = Font(name=FONT, size=10, color=INK)
        ws.cell(row=r, column=7).font = Font(name=FONT, size=10, color=INK)
        ws.cell(row=r, column=6).border = BORDER
        ws.cell(row=r, column=7).border = BORDER

    notes(ws, ref, [
        "・4M（人・機械・材料・方法）が基本の軸。まず4つ全部で切ってから、差が出た軸を深掘る",
        "・層別できないのは、記録時に層別キーを取っていないから。取得設計からやり直す",
        "・どの層でも差が出ないなら、キーの選び方が的外れ。別の軸を試す",
        "・件数（分母）が違う層を、件数のまま比べない。必ず率に直す",
        "・1水準10件未満の層は判断に使わない。偶然の差が大きく出る",
    ])

    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 40
    ch.add_data(Reference(ws, min_col=6, min_row=ref["first_row"], max_row=ref["last_row"]),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=3, min_row=ref["first_row"] + 1,
                                max_row=ref["last_row"]))
    solid(ch.series[0], NAVY)
    value_labels(ch)
    ch.legend = None
    ch.y_axis.scaling.orientation = "maxMin"
    ch.title = "層別ごとの不良率"
    ch.x_axis.title = "不良率(%)"
    ch.width, ch.height = 20, 14
    chart_font(ch)
    ws.add_chart(ch, "I5")
    return save(wb, outdir, "05_層別集計.xlsx")


def t_check_sheet(outdir):
    """記録用チェックシート。データを作る道具。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "チェックシート"
    header(ws, "チェックシート（データを作る・抜けを防ぐ）",
           "現場で記録を取るための用紙。層別キーを最初から列に持たせる。",
           "取得期間と件数を先に決めてから配る")

    rows = [["日付", "時刻", "号機", "作業者", "材料ロット", "品番",
             "不良項目", "件数", "備考"]]
    for _ in range(25):
        rows.append(["", "", "", "", "", "", "", "", ""])
    ref = table(ws, rows, input_cols=tuple(range(9)))
    for c, w in zip("BCDEFGHIJ", [12, 10, 10, 12, 14, 14, 20, 8, 24]):
        ws.column_dimensions[c].width = w

    notes(ws, ref, [
        "・層別キー（号機・作業者・ロット・時間帯）を最初から列にする。後からは作れない",
        "・不良項目は自由記述にせず、選択肢を決めて凡例に置く。表記ゆれで集計できなくなる",
        "・「その他」を選んだら備考に必ず内容を書く。書かせないと分類が育たない",
        "・記録する人が1回で書ける量にする。列が多すぎると記入されなくなる",
        "・取得期間・対象・除外条件を用紙の上に書いておく。後から条件が分からなくなる",
    ])

    # 凡例（不良項目の選択肢）
    r = ref["last_row"] + 9
    ws.cell(row=r, column=2, value="不良項目の選択肢（この中から選ぶ）").font = Font(
        name=FONT, size=10, bold=True, color=NAVY)
    for i, name in enumerate(["寸法不良（外径）", "キズ・打痕", "組付け位置ずれ",
                              "表面粗さ不良", "刻印かすれ", "その他（備考に記入）"], start=1):
        cell = ws.cell(row=r + i, column=2, value=f"{i}. {name}")
        cell.font = Font(name=FONT, size=9, color=INK)
    return save(wb, outdir, "06_チェックシート.xlsx")


BUILDERS = [t_pareto, t_histogram, t_control_chart, t_scatter_stratified,
            t_stratification, t_check_sheet]


def _assets_templates():
    """図テンプレートの置き場所。**このスキルの中。**

    スキルは `~/.claude/skills/<name>` から参照される。
    **その外に置いたものは、スキルの一部ではない。**
    実際に、リポジトリ直下に置いたところ
    「スキルだけを配ると雛形が0件」「規約の相対パスが空振り」になった。

    `__file__` は登録のリンク越しに呼ばれるので、**resolve() してから辿る。**
    """
    return pathlib.Path(__file__).resolve().parent.parent / "assets" / "templates"


def main():
    import argparse

    ap = argparse.ArgumentParser(description="QC7つ道具テンプレート（Excel）を生成する")
    ap.add_argument("outdir", nargs="?", default=None, help="出力先")
    ap.add_argument("--latin", default=None,
                    help="欧文フォント名。省くと日本語と同じものを使う")
    ap.add_argument("--font", default=FONT,
                    help="日本語フォント名。**資料テンプレートに合わせる。**"
                         "read_template_style.py で調べた値を渡す")
    args = ap.parse_args()

    # フォントは本スキルで固定しない。資料テンプレートが決める
    globals()["FONT"] = args.font
    globals()["LATIN"] = args.latin

    outdir = Path(args.outdir) if args.outdir else (
        _assets_templates() / "qc7"
    )
    outdir.mkdir(parents=True, exist_ok=True)
    print(f"QC7つ道具テンプレート（Excel）を生成: {outdir}  フォント: 日本語={FONT} / 欧文={LATIN or FONT}")
    for fn in BUILDERS:
        fn(outdir)
    print(f"完了: {len(BUILDERS)} 件")


if __name__ == "__main__":
    main()
